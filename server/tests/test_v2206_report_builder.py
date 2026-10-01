"""De Report Builder van Trends: inhoud, opmaak en uitvoer.

De directe knop "Exporteren naar PDF" is vervangen door een exportomgeving in
vijf stappen. De eisen die deze tests vastleggen:

* één databron: alle stappen rekenen met dezelfde gefilterde dataset;
* één grafieklogica: scherm, voorbeeld, PDF, Excel en PNG delen de renderer;
* wat je in het voorbeeld ziet is wat er uit de PDF komt;
* een onderdeel zonder gegevens doet niet mee, en een onderdeel dat struikelt
  neemt het rapport niet mee.
"""
import os
import sys
import unittest
from datetime import date, timedelta
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

ROOT = Path(__file__).resolve().parents[2]
from server.tests.desktop_source import desktop_source
sys.path.insert(0, str(ROOT))

from PySide6.QtWidgets import QApplication

import emt_report as report
from emt_charts import has_values, paint_trend_chart, render_trend_image
from emt_report_export import export_chart_png, export_excel, export_pdf
from emt_report_layout import ReportLayout
from emt_trends import METRICS, PERIODS, build_series

APP_SOURCE = desktop_source(ROOT)
UI_SOURCE = (ROOT / "emt_report_ui.py").read_text(encoding="utf-8")


def samenvatting(naam, dagen_terug, aangemeld=40, aanwezig=30, **extra):
    moment = date.today() - timedelta(days=dagen_terug)
    gegevens = {
        "id": naam.lower().replace(" ", "-"),
        "name": naam,
        "date": moment.strftime("%d-%m-%Y"),
        "event_type": extra.get("event_type", "Meeloopdag"),
        "place": extra.get("place", "Den Helder"),
        "location": extra.get("location", "Nieuwe Haven"),
        "source": "Eigen dossier",
        "statistiek": {
            "aangemeld": aangemeld, "aanwezig": aanwezig,
            "noshows": max(0, aangemeld - aanwezig),
            "opkomst_percentage": round(aanwezig / aangemeld * 100, 1) if aangemeld else 0.0,
            "noshow_percentage": round((aangemeld - aanwezig) / aangemeld * 100, 1) if aangemeld else 0.0,
            # De verdelingen zitten in de momentopname, net als in Trends zelf.
            "verdeling": extra.get("verdeling", {
                "Opleidingsniveau": {
                    "Mbo niveau 4 techniek": {"aangemeld": aangemeld - 10,
                                              "aanwezig": aanwezig - 8,
                                              "noshow": 2, "afgemeld": 1},
                    "Hbo bedrijfskunde": {"aangemeld": 10, "aanwezig": 8,
                                          "noshow": 1, "afgemeld": 1},
                },
                "Geslacht": {
                    "Man": {"aangemeld": aangemeld - 15, "aanwezig": aanwezig - 10,
                            "noshow": 3, "afgemeld": 2},
                    "Vrouw": {"aangemeld": 15, "aanwezig": 10, "noshow": 2, "afgemeld": 1},
                },
            }),
        },
    }
    return gegevens


def dataset(aantal=6):
    return [samenvatting(f"Evenement {nummer}", 30 * nummer, 40 + nummer * 5, 30 + nummer * 3)
            for nummer in range(1, aantal + 1)]


class ChartSpecTests(unittest.TestCase):
    """Elke grafiek in de builder is een analyse die Trends zelf ook maakt."""

    def test_every_chart_uses_a_known_metric(self):
        meetwaarden = {waarde for _label, waarde in METRICS}
        for sectie in report.SECTIONS:
            for grafiek in sectie["charts"]:
                with self.subTest(grafiek=grafiek["key"]):
                    self.assertIn(grafiek["metric"], meetwaarden)

    def test_the_time_unit_belongs_to_the_report_not_to_a_chart(self):
        """Een eigen vaste tijdseenheid per grafiek leverde één punt op.

        Een uitsplitsing naar groep stond op 'per jaar'; met gegevens uit één
        jaar bleef daar één meetmoment van over, en door één punt gaat geen
        lijn. Nu volgt elke grafiek de tijdseenheid van het rapport.
        """
        for sectie in report.SECTIONS:
            for grafiek in sectie["charts"]:
                with self.subTest(grafiek=grafiek["key"]):
                    self.assertNotIn("period", grafiek)
        self.assertEqual({waarde for _label, waarde in PERIODS}, set(report.PERIOD_LABELS))

    def test_chart_keys_are_unique(self):
        sleutels = [grafiek["key"] for sectie in report.SECTIONS for grafiek in sectie["charts"]]

        self.assertEqual(len(sleutels), len(set(sleutels)))

    def test_a_chart_spec_produces_a_real_series(self):
        gegevens = dataset()
        grafiek = report.SECTION_BY_KEY["opkomst"]["charts"][0]

        reeks = build_series(gegevens, grafiek["metric"], grafiek["dimension"], "month")

        self.assertTrue(has_values(reeks))


class EmptyPeriodTests(unittest.TestCase):
    """Niets invullen hoort te betekenen: alles wat bekend is.

    Met een lege periode wist het rapport niet welke periode het besloeg.
    Daardoor verdween de vergelijking en zei de titel niets, terwijl de
    gegevens er wel allemaal in zaten.
    """

    def setUp(self):
        self.gegevens = dataset(6)
        self.config = report.default_config("volledig")
        self.model = report.build_report_model(self.gegevens, self.config)

    def test_the_range_comes_from_the_oldest_and_newest_event(self):
        from emt_models import parse_date

        datums = sorted(parse_date(item["date"]) for item in self.gegevens)

        self.assertEqual(report.data_range(self.gegevens), (datums[0], datums[-1]))

    def test_an_empty_period_resolves_to_everything(self):
        eerste, laatste, afgeleid = report.effective_range(self.gegevens, {})

        self.assertTrue(afgeleid)
        self.assertEqual((eerste, laatste), report.data_range(self.gegevens))

    def test_a_chosen_period_is_left_alone(self):
        gekozen = {"since": date(2026, 1, 1), "until": date(2026, 6, 30)}

        eerste, laatste, afgeleid = report.effective_range(self.gegevens, gekozen)

        self.assertFalse(afgeleid)
        self.assertEqual((eerste, laatste), (gekozen["since"], gekozen["until"]))

    def test_only_one_date_filled_in_completes_the_other(self):
        _eerste, laatste, afgeleid = report.effective_range(
            self.gegevens, {"since": date(2026, 1, 1)}
        )

        self.assertTrue(afgeleid)
        self.assertEqual(laatste, report.data_range(self.gegevens)[1])

    def test_the_report_names_the_period_it_covers(self):
        self.assertIn("alles wat bekend is", self.model["periode"])
        self.assertNotEqual(self.model["titel"], "Trendsrapport")

    def test_step_one_shows_what_that_resolves_to(self):
        regels = dict(report.dataset_summary(self.model["summaries"], self.config["selectie"]))

        self.assertIn("alles wat bekend is", regels["Rapportperiode"])

    def test_the_development_section_appears_anyway(self):
        """Zonder gekozen periode is er geen vorige periode, wel een ontwikkeling."""
        sectie = next(item for item in self.model["sections"] if item["key"] == "vergelijking")

        self.assertEqual(sectie["title"], "Ontwikkeling binnen de periode")
        tabel = next(blok for blok in sectie["blocks"] if blok["kind"] == "table")
        self.assertEqual(tabel["columns"][1:3], ["Tweede helft", "Eerste helft"])

    def test_it_says_plainly_why_it_compares_halves(self):
        sectie = next(item for item in self.model["sections"] if item["key"] == "vergelijking")
        tekst = next(blok for blok in sectie["blocks"] if blok["kind"] == "paragraph")["text"]

        self.assertIn("geen voorgaande periode", tekst)

    def test_a_chosen_period_still_compares_with_the_previous_one(self):
        config = report.default_config("volledig")
        config["selectie"].update({"since": date.today() - timedelta(days=120),
                                   "until": date.today()})

        model = report.build_report_model(self.gegevens, config)
        sectie = next(item for item in model["sections"] if item["key"] == "vergelijking")

        self.assertEqual(sectie["title"], "Vergelijking met de vorige periode")
        tabel = next(blok for blok in sectie["blocks"] if blok["kind"] == "table")
        self.assertEqual(tabel["columns"][1:3], ["Deze periode", "Vorige periode"])

    def test_the_summary_wording_matches_what_is_compared(self):
        zin = report.management_summary(self.model["summaries"], self.model["vergelijking"])

        self.assertIn("tweede helft van de periode", zin)
        self.assertNotIn("voorgaande vergelijkbare periode", zin)

    def test_without_any_data_nothing_is_derived(self):
        self.assertEqual(report.data_range([]), (None, None))
        self.assertEqual(report.effective_range([], {})[:2], (None, None))


class BreakdownTests(unittest.TestCase):
    """Uitsplitsen naar de standen: waar zitten de no-shows en afmeldingen?"""

    def _table(self, sectie_sleutel="opleidingsniveau", gegevens=None):
        config = report.default_config("volledig")
        model = report.build_report_model(gegevens or dataset(), config)
        sectie = next(item for item in model["sections"] if item["key"] == sectie_sleutel)
        return next(blok for blok in sectie["blocks"] if blok["kind"] == "table")

    def test_the_table_shows_every_state_per_group(self):
        tabel = self._table()

        self.assertEqual(tabel["columns"],
                         ["Opleidingsniveau", "Aanmeldingen", "Aandeel", "Aanwezig",
                          "No-shows", "Afmeldingen", "Opkomst"])

    def test_the_numbers_add_up_per_group(self):
        tabel = self._table()
        rij = next(regel for regel in tabel["rows"] if regel[0].startswith("Mbo"))

        # Zes evenementen met elk twee no-shows en één afmelding in deze groep.
        self.assertEqual(rij[4], "12")
        self.assertEqual(rij[5], "6")

    def test_every_group_dimension_can_be_split(self):
        for sleutel in ("opleidingsniveau", "geslacht"):
            with self.subTest(sectie=sleutel):
                self.assertGreater(len(self._table(sleutel)["rows"]), 1)

    def test_a_chart_exists_for_noshows_and_cancellations(self):
        for sectie in ("leeftijd", "geslacht", "opleidingsniveau", "profiel"):
            sleutels = {grafiek["key"] for grafiek in report.SECTION_BY_KEY[sectie]["charts"]}
            with self.subTest(sectie=sectie):
                self.assertTrue(any(sleutel.endswith("_noshows") for sleutel in sleutels))
                self.assertTrue(any(sleutel.endswith("_afmeldingen") for sleutel in sleutels))

    def test_an_older_snapshot_leaves_the_cancellation_cell_empty(self):
        """Vóór schema 4 waren afmeldingen per groep niet vastgelegd."""
        oud = []
        for item in dataset(2):
            verdeling = {
                dimensie: {naam: {"aangemeld": waarden["aangemeld"],
                                  "aanwezig": waarden["aanwezig"]}
                           for naam, waarden in groepen.items()}
                for dimensie, groepen in item["statistiek"]["verdeling"].items()
            }
            oud.append(dict(item, statistiek=dict(item["statistiek"], verdeling=verdeling)))

        tabel = self._table(gegevens=oud)

        # Zonder vastgelegde afmeldingen per groep blijft die cel leeg in
        # plaats van ten onrechte nul te tonen.
        self.assertTrue(all(regel[5] == "—" for regel in tabel["rows"]))

    def test_cancellations_are_a_metric_of_their_own(self):
        from emt_trends import METRICS

        self.assertIn(("Afmeldingen", "afgemeld"), METRICS)

    def test_the_key_figures_mention_cancellations(self):
        model = report.build_report_model(dataset(), report.default_config("volledig"))
        kpis = next(blok for sectie in model["sections"]
                    for blok in sectie["blocks"] if blok["kind"] == "kpis")

        self.assertIn("Afmeldingen", [label for label, _waarde in kpis["items"]])


class PeriodTests(unittest.TestCase):
    """De tijdseenheid van de ontwikkelingsgrafieken."""

    def _chart(self, periode, sleutel="opkomst_tijd"):
        config = report.default_config("volledig")
        config["selectie"]["periode"] = periode
        model = report.build_report_model(dataset(8), config)
        return next(blok for sectie in model["sections"] for blok in sectie["blocks"]
                    if blok["kind"] == "chart" and blok.get("chart_key") == sleutel)

    def test_every_trend_period_is_offered(self):
        self.assertEqual(set(report.PERIOD_LABELS),
                         {"event", "month", "quarter", "year"})

    def test_a_coarser_unit_gives_fewer_points(self):
        per_maand = self._chart("month")["series"]["points"]
        per_jaar = self._chart("year")["series"]["points"]

        self.assertGreater(len(per_maand), len(per_jaar))

    def test_per_event_labels_are_event_names(self):
        punten = self._chart("event")["series"]["points"]

        self.assertTrue(any("Evenement" in str(punt["label"]) for punt in punten))

    def test_a_group_breakdown_is_a_development_too(self):
        """Hier zat de fout: per groep bleef één punt over, dus alleen bolletjes."""
        reeks = self._chart("month", "niveau_noshows")["series"]

        self.assertGreater(len(reeks["points"]), 1)
        self.assertGreater(len(reeks["groups"]), 1)

    def test_an_unknown_unit_falls_back_to_months(self):
        self.assertEqual(report.selected_period({"periode": "onzin"}), "month")
        self.assertEqual(report.selected_period({}), "month")

    def test_the_dataset_summary_names_the_unit(self):
        regels = dict(report.dataset_summary(dataset(), {"periode": "quarter"}))

        self.assertEqual(regels["Tijdseenheid"], "Per kwartaal")

    def test_percentage_choice_is_carried_into_charts_and_tables(self):
        config = report.default_config("volledig")
        config["selectie"]["percentage_of_total"] = True
        model = report.build_report_model(dataset(2), config)
        blocks = [block for section in model["sections"] for block in section["blocks"]]
        percentage_charts = [
            block for block in blocks
            if block["kind"] == "chart" and block["series"].get("percentage_of_total")
        ]
        self.assertTrue(percentage_charts)
        tables = [block for block in blocks if block["kind"] == "table"]
        self.assertTrue(any("%" in str(cell) and "(" in str(cell)
                            for block in tables for row in block["rows"] for cell in row))


class ConfigTests(unittest.TestCase):
    def test_presets_only_configure_the_same_builder(self):
        """Geen aparte rapportimplementatie per profiel."""
        for sleutel, _label, keys in report.PRESETS:
            with self.subTest(preset=sleutel):
                onderdelen = report.sections_for_preset(sleutel)
                self.assertEqual([item["key"] for item in onderdelen],
                                 [sectie["key"] for sectie in report.SECTIONS])
                self.assertEqual({item["key"] for item in onderdelen if item["aan"]}, set(keys))

    def test_an_older_template_gains_new_sections_switched_off(self):
        oud = {"secties": [{"key": "kerncijfers", "aan": True}], "vormgeving": {}}

        hersteld = report.prepare_config(oud)

        self.assertEqual([item["key"] for item in hersteld["secties"]][0], "kerncijfers")
        self.assertEqual(len(hersteld["secties"]), len(report.SECTIONS))
        self.assertTrue(hersteld["secties"][0]["aan"])
        self.assertFalse(any(item["aan"] for item in hersteld["secties"][1:]))

    def test_a_template_keeps_the_shape_and_not_the_data(self):
        config = report.default_config("opkomst")
        config["selectie"]["since"] = date(2026, 1, 1)

        sjabloon = report.template_payload("Kwartaalrapport Werving", config)

        self.assertEqual(sjabloon["naam"], "Kwartaalrapport Werving")
        self.assertIn("secties", sjabloon)
        self.assertIn("vormgeving", sjabloon)
        self.assertNotIn("selectie", sjabloon)
        self.assertNotIn("summaries", sjabloon)

    def test_a_template_can_be_reused_on_another_period(self):
        sjabloon = report.template_payload("Kwartaal", report.default_config("opkomst"))
        nieuw = {"since": date(2026, 4, 1), "until": date(2026, 6, 30)}

        config = report.config_from_template(sjabloon, nieuw)

        self.assertEqual(config["selectie"]["since"], date(2026, 4, 1))
        self.assertEqual({item["key"] for item in config["secties"] if item["aan"]},
                         set(report.PRESET_BY_KEY["opkomst"]))

    def test_the_order_of_the_sections_is_the_order_of_the_report(self):
        config = report.default_config("volledig")
        config["secties"].reverse()

        model = report.build_report_model(dataset(), config)
        volgorde = [sectie["key"] for sectie in model["sections"]]

        gewenst = [item["key"] for item in config["secties"]
                   if item["aan"] and item["key"] in volgorde]
        self.assertEqual(volgorde, gewenst)


class FileNameTests(unittest.TestCase):
    def test_forbidden_characters_are_removed(self):
        self.assertEqual(report.safe_file_name('Rapport: Q2 <intern>?'), "Rapport Q2 intern")

    def test_an_empty_name_falls_back(self):
        self.assertEqual(report.safe_file_name("   "), "Trendsrapport")
        self.assertEqual(report.safe_file_name("..."), "Trendsrapport")

    def test_reserved_windows_names_fall_back(self):
        self.assertEqual(report.safe_file_name("CON"), "Trendsrapport")


class AvailabilityTests(unittest.TestCase):
    """Onderdelen zonder bruikbare gegevens doen niet mee."""

    def test_without_data_nothing_is_available(self):
        self.assertEqual(report.available_sections([]), set())

    def test_missing_distributions_hide_their_section(self):
        zonder = [dict(item, statistiek=dict(item["statistiek"], verdeling={}))
                  for item in dataset()]

        beschikbaar = report.available_sections(zonder)

        self.assertNotIn("opleidingsniveau", beschikbaar)
        self.assertNotIn("geslacht", beschikbaar)
        self.assertIn("kerncijfers", beschikbaar)

    def test_present_distributions_enable_their_section(self):
        beschikbaar = report.available_sections(dataset())

        self.assertIn("opleidingsniveau", beschikbaar)
        self.assertIn("geslacht", beschikbaar)

    def test_charts_without_values_are_reported_as_unavailable(self):
        zonder = [dict(item, statistiek=dict(item["statistiek"], verdeling={}))
                  for item in dataset()]

        self.assertEqual(report.available_charts(zonder, "opleidingsniveau"), set())
        self.assertTrue(report.available_charts(dataset(), "opleidingsniveau"))


class SummaryTextTests(unittest.TestCase):
    """De managementsamenvatting komt uit sjablonen, niet uit een taalmodel."""

    def test_the_numbers_come_from_the_data(self):
        tekst = report.management_summary(dataset(2))

        self.assertIn("2 evenementen", tekst)
        self.assertIn("aanmeldingen", tekst)
        self.assertIn("opkomstpercentage", tekst)

    def test_it_is_reproducible(self):
        gegevens = dataset()

        self.assertEqual(report.management_summary(gegevens),
                         report.management_summary(gegevens))

    def test_it_does_not_claim_causality(self):
        tekst = report.management_summary(dataset())

        self.assertNotIn("doordat", tekst.lower())
        self.assertNotIn("veroorzaakt", tekst.lower())

    def test_one_event_reads_as_one_event(self):
        self.assertIn("is 1 evenement", report.management_summary(dataset(1)))

    def test_without_data_it_says_so(self):
        self.assertIn("geen evenementen", report.management_summary([]))


class LayoutTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.qt = QApplication.instance() or QApplication([])

    def _layout(self, preset="volledig", **vormgeving):
        config = report.default_config(preset)
        config["vormgeving"].update(vormgeving)
        return ReportLayout(report.build_report_model(dataset(8), config))

    def test_a_report_has_pages(self):
        self.assertGreater(self._layout().page_count(), 1)

    def test_the_cover_can_be_switched_off(self):
        met = self._layout(voorblad=True)
        zonder = self._layout(voorblad=False)

        self.assertTrue(met.pages[0]["cover"])
        self.assertFalse(zonder.pages[0]["cover"])

    def test_a_chart_never_lands_half_on_a_page(self):
        """Een grafiek is één blok; die wordt nooit doorgesneden."""
        opmaak = self._layout()
        for pagina in opmaak.pages:
            for _x, y, _breedte, hoogte, blok in pagina["items"]:
                if blok["kind"] == "chart":
                    with self.subTest(titel=blok.get("title")):
                        self.assertLessEqual(y + hoogte, opmaak.content_bottom + 1)

    def test_a_title_is_never_alone_at_the_bottom(self):
        """Een kop zonder inhoud eronder maakt een pagina rommelig."""
        opmaak = self._layout()
        for nummer, pagina in enumerate(opmaak.pages):
            if not pagina["items"]:
                continue
            laatste = pagina["items"][-1][4]
            with self.subTest(pagina=nummer + 1):
                self.assertNotIn(laatste["kind"], {"heading", "subheading"})

    def test_a_long_table_repeats_its_header(self):
        config = report.default_config("aangepast")
        for onderdeel in config["secties"]:
            onderdeel["aan"] = onderdeel["key"] == "evenementen"
        model = report.build_report_model(dataset(60), config)
        opmaak = ReportLayout(model)

        stukken = [blok for pagina in opmaak.pages for *_r, blok in pagina["items"]
                   if blok["kind"] == "table"]

        self.assertGreater(len(stukken), 1, "zo'n lange tabel hoort te breken")
        for stuk in stukken:
            self.assertEqual(stuk["columns"][0], "Evenement")

    def test_orientation_follows_the_choice(self):
        self.assertEqual(self._layout(orientatie="liggend").orientation, "liggend")
        self.assertEqual(self._layout(orientatie="staand").orientation, "staand")

    def test_landscape_pages_are_wider_than_tall(self):
        opmaak = self._layout(orientatie="liggend")

        self.assertGreater(opmaak.page_width, opmaak.page_height)

    def test_page_numbers_can_be_switched_off(self):
        self.assertTrue(self._layout(paginanummers=True)._has_footer())

    def test_an_empty_report_still_has_a_page(self):
        config = report.default_config("aangepast")
        opmaak = ReportLayout(report.build_report_model([], config))

        self.assertGreaterEqual(opmaak.page_count(), 1)


class DisplayChoiceTests(unittest.TestCase):
    """Statistieken of verloop: dezelfde cijfers, twee manieren van kijken."""

    def _charts(self, weergave):
        config = report.default_config("volledig")
        config["selectie"]["weergave"] = weergave
        model = report.build_report_model(dataset(6), config)
        return [blok for sectie in model["sections"] for blok in sectie["blocks"]
                if blok["kind"] == "chart"]

    def test_both_ways_are_offered(self):
        self.assertEqual(set(report.DISPLAY_LABELS), {report.STATISTIEKEN, report.VERLOOP})

    def test_development_spreads_the_data_over_time(self):
        for grafiek in self._charts(report.VERLOOP):
            with self.subTest(grafiek=grafiek["chart_key"]):
                self.assertGreater(len(grafiek["series"]["points"]), 1)

    def test_statistics_put_the_whole_selection_in_one_picture(self):
        """Eén punt per groep: dat tekent als staven in plaats van lijnen."""
        for grafiek in self._charts(report.STATISTIEKEN):
            with self.subTest(grafiek=grafiek["chart_key"]):
                self.assertEqual(len(grafiek["series"]["points"]), 1)
                self.assertEqual(grafiek["series"]["points"][0]["label"], "Hele selectie")

    def test_the_numbers_are_the_same_either_way(self):
        """Een andere weergave mag de cijfers niet veranderen."""
        def totaal(grafieken, sleutel):
            grafiek = next(item for item in grafieken if item["chart_key"] == sleutel)
            return sum(sum(punt["values"].values()) for punt in grafiek["series"]["points"])

        self.assertAlmostEqual(totaal(self._charts(report.VERLOOP), "niveau_aanmeldingen"),
                               totaal(self._charts(report.STATISTIEKEN), "niveau_aanmeldingen"),
                               delta=0.1)

    def test_statistics_hide_the_time_unit(self):
        """Zonder tijdas zegt een tijdseenheid niets meer."""
        regels = dict(report.dataset_summary(dataset(6), {"weergave": report.STATISTIEKEN}))

        self.assertEqual(regels["Weergave"], report.DISPLAY_LABELS[report.STATISTIEKEN])
        self.assertNotIn("Tijdseenheid", regels)

    def test_development_names_the_time_unit(self):
        regels = dict(report.dataset_summary(
            dataset(6), {"weergave": report.VERLOOP, "periode": "quarter"}
        ))

        self.assertEqual(regels["Tijdseenheid"], "Per kwartaal")

    def test_an_unknown_choice_falls_back_to_development(self):
        self.assertEqual(report.selected_display({"weergave": "onzin"}), report.VERLOOP)
        self.assertEqual(report.selected_display({}), report.VERLOOP)

    def test_the_builder_offers_the_choice(self):
        self.assertIn("self.display_choice = QComboBox()", UI_SOURCE)
        self.assertIn('vorm.addRow("Weergave:", self.display_choice)', UI_SOURCE)
        # De tijdseenheid gaat op slot zodra er geen tijdas meer is.
        self.assertIn("self.period_choice.setEnabled(verloop)", UI_SOURCE)

    def test_a_deliberate_single_moment_needs_no_excuse(self):
        """De melding over te weinig meetmomenten hoort hier niet."""
        bron = (ROOT / "emt_charts.py").read_text(encoding="utf-8")

        self.assertIn('if series.get("period") != "total":', bron)


class ReportStructureTests(unittest.TestCase):
    """Eén tabel per soort, en elk onderdeel begint op een eigen pagina."""

    @classmethod
    def setUpClass(cls):
        cls.qt = QApplication.instance() or QApplication([])

    def _model(self, preset="volledig"):
        return report.build_report_model(dataset(6), report.default_config(preset))

    def test_no_table_appears_twice(self):
        """Detailtabellen herhaalde precies wat de hoofdstukken al toonden."""
        model = self._model()
        koppen = [tuple(blok["columns"]) for sectie in model["sections"]
                  for blok in sectie["blocks"] if blok["kind"] == "table"]

        self.assertEqual(len(koppen), len(set(koppen)))

    def test_a_dimension_with_its_own_chapter_is_skipped_in_the_detail_section(self):
        model = self._model()
        secties = {sectie["key"] for sectie in model["sections"]}

        # Alle uitsplitsingen hebben een eigen hoofdstuk, dus houdt de aparte
        # sectie niets meer over en verdwijnt hij.
        self.assertIn("opleidingsniveau", secties)
        self.assertNotIn("detailtabellen", secties)

    def test_the_detail_section_still_works_on_its_own(self):
        """Wie alleen de tabellen wil, krijgt ze nog steeds."""
        config = report.default_config("aangepast")
        for onderdeel in config["secties"]:
            onderdeel["aan"] = onderdeel["key"] == "detailtabellen"

        model = report.build_report_model(dataset(6), config)
        sectie = next(item for item in model["sections"] if item["key"] == "detailtabellen")

        self.assertTrue([blok for blok in sectie["blocks"] if blok["kind"] == "table"])

    def test_every_section_starts_on_its_own_page(self):
        """Twee hoofdstukken op één blad laten de lezer de grens zoeken."""
        opmaak = ReportLayout(self._model())

        for nummer, pagina in enumerate(opmaak.pages):
            koppen = [blok for *_rest, blok in pagina["items"] if blok["kind"] == "heading"]
            with self.subTest(pagina=nummer + 1):
                self.assertLessEqual(len(koppen), 1)

    def test_a_heading_is_always_the_first_thing_on_its_page(self):
        opmaak = ReportLayout(self._model())

        for nummer, pagina in enumerate(opmaak.pages):
            if not pagina["items"]:
                continue
            soorten = [blok["kind"] for *_rest, blok in pagina["items"]]
            if "heading" in soorten:
                with self.subTest(pagina=nummer + 1):
                    self.assertEqual(soorten.index("heading"), 0)


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.qt = QApplication.instance() or QApplication([])
        cls.model = report.build_report_model(dataset(8), report.default_config("volledig"))

    def setUp(self):
        import tempfile

        self.map = Path(tempfile.mkdtemp(prefix="eventhub-test-"))

    def test_the_pdf_is_a_real_pdf(self):
        doel = export_pdf(self.model, self.map / "rapport.pdf")

        self.assertTrue(doel.exists())
        self.assertEqual(doel.read_bytes()[:4], b"%PDF")

    def test_the_pdf_has_the_same_pages_as_the_preview(self):
        """Het voorbeeld en de PDF gebruiken dezelfde opmaak."""
        opmaak = ReportLayout(self.model)
        export_pdf(self.model, self.map / "rapport.pdf")

        self.assertGreater(opmaak.page_count(), 0)
        self.assertIn("ReportLayout(", (ROOT / "emt_report_export.py").read_text(encoding="utf-8"))
        self.assertIn("ReportLayout(", UI_SOURCE)

    def test_the_workbook_has_a_dashboard_and_data_sheets(self):
        from openpyxl import load_workbook

        doel = export_excel(self.model, self.map / "rapport.xlsx")
        werkmap = load_workbook(doel)

        self.assertEqual(werkmap.sheetnames[0], "Samenvatting")
        self.assertIn("Evenementen", werkmap.sheetnames)
        self.assertGreater(len(werkmap.sheetnames), 2)

    def test_the_dashboard_carries_the_eventhub_charts(self):
        from openpyxl import load_workbook

        doel = export_excel(self.model, self.map / "rapport.xlsx")
        dashboard = load_workbook(doel)["Samenvatting"]

        self.assertTrue(dashboard._images, "het dashboard hoort de gerenderde grafieken te bevatten")

    def test_the_data_sheets_use_real_types(self):
        from openpyxl import load_workbook

        doel = export_excel(self.model, self.map / "rapport.xlsx")
        blad = load_workbook(doel)["Evenementen"]

        self.assertEqual(blad.freeze_panes, "A2")
        self.assertTrue(blad.auto_filter.ref)
        self.assertIsInstance(blad["B2"].value, date)
        self.assertIn("%", blad["H2"].number_format)
        self.assertIsInstance(blad["E2"].value, int)

    def test_a_chart_exports_as_a_sharp_png(self):
        grafiek = next(blok for sectie in self.model["sections"]
                       for blok in sectie["blocks"] if blok["kind"] == "chart")

        doel = export_chart_png(grafiek["series"], self.map / "grafiek.png", 800, 300, scale=2.0)

        self.assertTrue(doel.exists())
        from PySide6.QtGui import QImage

        beeld = QImage(str(doel))
        self.assertEqual((beeld.width(), beeld.height()), (1600, 600))

    def test_the_png_uses_the_same_renderer_as_the_screen(self):
        bron = (ROOT / "emt_report_export.py").read_text(encoding="utf-8")

        self.assertIn("render_trend_image(", bron)
        self.assertIn("def render_trend_image(", (ROOT / "emt_charts.py").read_text(encoding="utf-8"))


class RobustnessTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.qt = QApplication.instance() or QApplication([])

    def test_a_broken_section_does_not_take_the_report_down(self):
        model = report.build_report_model(dataset(), report.default_config("volledig"))
        model["sections"][0]["blocks"].insert(0, {"kind": "table", "columns": ["A"],
                                                  "rows": [[None]], "weights": None})

        opmaak = ReportLayout(model)

        self.assertGreater(opmaak.page_count(), 0)

    def test_a_chart_without_values_is_skipped(self):
        config = report.default_config("volledig")
        model = report.build_report_model([], config)

        for sectie in model["sections"]:
            for blok in sectie["blocks"]:
                self.assertNotEqual(blok["kind"], "chart")

    def test_an_empty_series_still_draws_something(self):
        beeld = render_trend_image({"groups": [], "points": []}, 300, 120)

        self.assertEqual((beeld.width(), beeld.height()), (600, 240))


class ScreenTests(unittest.TestCase):
    """Wat er in Trends zelf is veranderd."""

    def test_the_button_opens_the_report_builder(self):
        self.assertIn('QPushButton("Exporteren")', APP_SOURCE)
        start = APP_SOURCE.index("def export_trend_data(self")
        blok = APP_SOURCE[start:APP_SOURCE.index("\n    def ", start + 1)]
        self.assertIn("ReportBuilderDialog(", blok)

    def test_the_old_direct_pdf_flow_is_gone(self):
        self.assertNotIn("_trend_pdf_document", APP_SOURCE)
        self.assertNotIn('QUrl("trend://chart")', APP_SOURCE)

    def test_the_builder_has_five_steps(self):
        from emt_report_ui import STAPPEN

        self.assertEqual(STAPPEN, ["Selectie", "Inhoud", "Vormgeving",
                                   "Exportvoorbeeld", "Exporteren"])

    def test_the_builder_starts_from_the_active_filters(self):
        self.assertIn("panel.current_filters()", APP_SOURCE)
        start = APP_SOURCE.index("def current_filters(self")
        blok = APP_SOURCE[start:APP_SOURCE.index("\n    def ", start + 1)]
        for sleutel in ("since", "event_type", "location", "event_id"):
            self.assertIn(f'"{sleutel}"', blok)

    def test_sections_can_be_reordered_by_dragging(self):
        self.assertIn("QAbstractItemView.DragDropMode.InternalMove", UI_SOURCE)
        self.assertIn("rowsMoved.connect", UI_SOURCE)

    def test_templates_can_be_saved_renamed_and_deleted(self):
        for handler in ("_save_template", "_rename_template", "_delete_template",
                        "_load_template"):
            with self.subTest(handler=handler):
                self.assertIn(f"def {handler}(", UI_SOURCE)

    def test_the_builder_offers_the_time_unit(self):
        self.assertIn("self.period_choice = QComboBox()", UI_SOURCE)
        self.assertIn('"periode": str(self.period_choice.currentData()', UI_SOURCE)

    def test_the_dataset_summary_uses_two_columns(self):
        """Onder elkaar liep de opsomming onderaan stap 1 buiten beeld."""
        self.assertIn("self.dataset_layout = QGridLayout(self.dataset_box)", UI_SOURCE)
        self.assertIn("DATASET_KOLOMMEN = 2", UI_SOURCE)
        # Kolomsgewijs vullen houdt de leesvolgorde van boven naar beneden.
        self.assertIn("index % per_kolom, index // per_kolom", UI_SOURCE)

    def test_the_preview_offers_paging_and_zoom(self):
        for onderdeel in ("_step_page", "_zoom", "_fit_preview", "Terug naar inhoud",
                          "Terug naar vormgeving"):
            with self.subTest(onderdeel=onderdeel):
                self.assertIn(onderdeel, UI_SOURCE)

    def test_a_single_chart_can_be_exported_as_png(self):
        self.assertIn("def export_png(self):", APP_SOURCE)
        self.assertIn("Exporteren als PNG", APP_SOURCE)
        self.assertIn("export_chart_png(", APP_SOURCE)

    def test_the_screen_chart_and_the_export_share_one_painter(self):
        self.assertIn("paint_trend_chart(", APP_SOURCE)
        self.assertIn("paint_trend_chart(", (ROOT / "emt_report_layout.py").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
