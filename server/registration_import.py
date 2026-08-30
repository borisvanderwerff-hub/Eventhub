"""Excel bezoekerslijst parsing for EventHub Server.

This is a deliberate, self-contained copy of the column-detection /
name / date / introducee logic that also exists in EventHub Desktop's
``bezoekerslijst_core.py``. It is duplicated on purpose: EventHub Server
must be runnable on a machine that has never had EventHub Desktop
installed, so nothing under ``server/`` may import from outside this
folder. If EventHub Desktop's import logic changes, re-sync this file
by hand (or, later, extract both into a small pip-installable shared
package that each product depends on explicitly).
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
import xlrd

STRING_FIELDS = [
    "Evenement", "Identifier", "GastVan", "Voornaam", "Tussenvoegsel", "Achternaam",
    "Email", "Telefoonnummer", "Geboortedatum", "Geboorteplaats", "Geslacht",
    "Opleiding", "Profiel", "Type", "Aanwezigheid", "Gebruik",
]

FIELD_ALIASES = {
    "Evenement": [
        "Evenement", "Event", "Naam evenement", "Evenementnaam", "Eventnaam",
        "Titel evenement", "Evenement titel", "Event title", "Eventtitel",
        "Naam event", "Activiteit", "Activiteitsnaam", "Bijeenkomst",
    ],
    "Identifier": ["Identifier", "ID", "Kandidaatnummer", "Registratienummer", "Registratie ID", "Registratie-ID"],
    "GastVan": ["Gast van", "Gastvan", "Guest of", "Introducé van", "Introducee van"],
    "Voornaam": ["Voornaam", "First name", "Firstname"],
    "Tussenvoegsel": ["Tussenvoegsel", "Prefix"],
    "Achternaam": ["Achternaam", "Last name", "Lastname"],
    "Email": ["E-mail", "Email", "E-mailadres"],
    "Telefoonnummer": ["Telefoonnummer", "Telefoon", "Mobiel", "Mobiele nummer", "Phone"],
    "Geboortedatum": ["Geboortedatum", "Date of birth", "Birth date"],
    "Geboorteplaats": [
        "Geboorteplaats", "Geboorte plaats", "Plaats van geboorte",
        "Plaats geboorte", "Geboortestad", "Geb. plaats", "Geb plaats",
        "Geboorteplaats volgens identiteitsbewijs", "Geboorteplaats ID",
        "Place of birth", "Birth place", "Birthplace", "Birth city",
        "City of birth", "Birth location",
    ],
    "Geslacht": ["Geslacht", "Gender"],
    "Opleiding": ["Opleiding", "Opleidingsniveau", "Niveau", "Education level"],
    "Profiel": [
        "Opleidingsrichting", "Profiel", "Studierichting",
        "Opleidingsprofiel", "Richting", "Study profile",
    ],
    "Type": ["Type"],
    "Aanwezigheid": [
        "Aanwezigheid", "Aanwezig", "Present", "Presence", "Attendance",
        "Opgekomen", "Show", "No show", "Noshow",
    ],
    "Gebruik": ["Gebruik"],
}

PRESENT_VALUES = {
    "ja", "yes", "y", "aanwezig", "present", "1", "true", "waar", "opgekomen", "gekomen", "show",
}
ABSENT_VALUES = {
    "nee", "no", "n", "afwezig", "absent", "0", "false", "onwaar",
    "noshow", "no-show", "nietopgekomen", "nietgekomen",
}

HEADER_PATTERNS = {
    "Evenement": ("evenement", "eventnaam", "eventtitel", "eventtitle", "naamevent", "activiteitsnaam"),
    "Geboorteplaats": (
        "geboorteplaats", "plaatsvangeboorte", "plaatsgeboorte", "geboortestad",
        "gebplaats", "placeofbirth", "birthplace", "birthcity", "cityofbirth", "birthlocation",
    ),
}


def normalize(value) -> str:
    value = unicodedata.normalize("NFD", str(value or "").strip().lower())
    value = "".join(char for char in value if unicodedata.category(char) != "Mn")
    return re.sub(r"[^a-z0-9]", "", value)


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def date_text(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).strftime("%d-%m-%Y")
        except (ValueError, OverflowError, TypeError):
            pass
    return text(value)


def phone_text(value) -> str:
    result = text(value)
    if result.endswith(".0") and result[:-2].isdigit():
        result = result[:-2]
    if len(result) == 9 and result.isdigit() and not result.startswith("0"):
        result = f"0{result}"
    return result


def infer_presence_from_text(value) -> bool | None:
    key = normalize(value)
    if not key:
        return None
    key = key.replace("-", "")
    if key in PRESENT_VALUES:
        return True
    if key in ABSENT_VALUES:
        return False
    return None


def is_introducee(record: dict) -> bool:
    return bool(text(record.get("GastVan", "")))


def _first_column(value) -> int:
    if isinstance(value, list):
        return value[0] if value else -1
    return int(value)


def build_map(headers: dict[str, list[int] | int]) -> dict[str, int]:
    result = {}
    for field, aliases in FIELD_ALIASES.items():
        result[field] = -1
        for alias in aliases:
            key = normalize(alias)
            if key in headers:
                result[field] = _first_column(headers[key])
                break
        if result[field] < 0:
            patterns = HEADER_PATTERNS.get(field, ())
            for header, columns in headers.items():
                if any(pattern in header for pattern in patterns):
                    result[field] = _first_column(columns)
                    break
    return result


def _looks_like_date_value(value) -> bool:
    if isinstance(value, (date, datetime)):
        return True
    if isinstance(value, (int, float)):
        return 1 <= value <= 100000
    value = text(value)
    if not value:
        return False
    patterns = (
        r"^\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}$",
        r"^\d{4}[-/.]\d{1,2}[-/.]\d{1,2}$",
    )
    return any(re.match(pattern, value) for pattern in patterns)


def _duplicate_birthplace_column(headers, rows, header_row_index, mapping) -> int:
    if mapping.get("Geboorteplaats", -1) >= 0:
        return -1
    birthdate_columns = []
    for alias in FIELD_ALIASES["Geboortedatum"]:
        columns = headers.get(normalize(alias), [])
        if not isinstance(columns, list):
            columns = [columns]
        birthdate_columns.extend(columns)
    birthdate_columns = sorted(set(birthdate_columns))
    if len(birthdate_columns) < 2:
        return -1

    first_birthdate = mapping.get("Geboortedatum", birthdate_columns[0])
    candidates = [column for column in birthdate_columns if column != first_birthdate]
    candidates.sort(key=lambda column: (column != first_birthdate + 1, column))
    for column in candidates:
        values = []
        for row in rows[header_row_index + 1:header_row_index + 41]:
            if column < len(row) and text(row[column]):
                values.append(row[column])
        if not values:
            return column
        date_values = sum(_looks_like_date_value(value) for value in values)
        if date_values / len(values) < 0.5:
            return column
    return -1


def _xlsx_sheets(file_path: Path):
    workbook = load_workbook(file_path, read_only=True, data_only=True, keep_vba=False)
    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            yield worksheet.title, rows
    finally:
        workbook.close()


def _xls_sheets(file_path: Path):
    workbook = xlrd.open_workbook(file_path)
    for worksheet in workbook.sheets():
        rows = []
        for row_index in range(worksheet.nrows):
            row = []
            for column_index in range(worksheet.ncols):
                cell = worksheet.cell(row_index, column_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = datetime(*xlrd.xldate_as_tuple(value, workbook.datemode))
                row.append(value)
            rows.append(row)
        yield worksheet.name, rows


def read_registration_file(file_path: str | Path):
    file_path = Path(file_path)
    sheets = list(_xls_sheets(file_path) if file_path.suffix.lower() == ".xls" else _xlsx_sheets(file_path))
    best = None
    for sheet_index, (sheet_name, rows) in enumerate(sheets):
        for row_index, row in enumerate(rows[:25]):
            headers = {}
            for column_index, value in enumerate(row):
                key = normalize(value)
                if key:
                    headers.setdefault(key, []).append(column_index)
            mapping = build_map(headers)
            if mapping["Voornaam"] < 0 or mapping["Achternaam"] < 0:
                continue
            repaired_birthplace_column = _duplicate_birthplace_column(headers, rows, row_index, mapping)
            if repaired_birthplace_column >= 0:
                mapping["Geboorteplaats"] = repaired_birthplace_column
            score = sum(index >= 0 for index in mapping.values())
            if best is None or score > best["score"]:
                best = {
                    "sheet_index": sheet_index,
                    "sheet_name": sheet_name,
                    "row_index": row_index,
                    "rows": rows,
                    "mapping": mapping,
                    "score": score,
                    "repaired_birthplace_header": repaired_birthplace_column >= 0,
                }
    if best is None:
        raise ValueError("Geen werkblad met de kolommen Voornaam en Achternaam gevonden.")

    records = []
    fallback_events = 0
    missing_birth_places = 0
    presence_detected = 0
    fallback_event = file_path.stem
    mapping = best["mapping"]

    for source in best["rows"][best["row_index"] + 1:]:
        def get(field):
            column = mapping[field]
            return source[column] if 0 <= column < len(source) else ""

        first_name = text(get("Voornaam"))
        last_name = text(get("Achternaam"))
        if not first_name and not last_name:
            continue
        record = {field: text(get(field)) for field in STRING_FIELDS}
        record.update({
            "Voornaam": first_name,
            "Achternaam": last_name,
            "Geboortedatum": date_text(get("Geboortedatum")),
            "Telefoonnummer": phone_text(get("Telefoonnummer")),
            "Aanwezig": False,
        })
        detected_presence = infer_presence_from_text(record.get("Aanwezigheid", ""))
        if detected_presence is not None:
            record["Aanwezig"] = detected_presence
            presence_detected += 1
        if not record["Evenement"]:
            record["Evenement"] = fallback_event
            fallback_events += 1
        if not record["Geboorteplaats"]:
            missing_birth_places += 1
        records.append(record)

    return {
        "records": records,
        "report": {
            "file_name": file_path.name,
            "sheet_name": best["sheet_name"],
            "introducees": sum(is_introducee(record) for record in records),
            "fallback_events": fallback_events,
            "missing_birth_places": missing_birth_places,
            "presence_detected": presence_detected,
            "repaired_birthplace_header": best.get("repaired_birthplace_header", False),
        },
    }
