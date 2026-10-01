"""De uitvoer van een rapport: PDF, Excel en losse grafieken.

Alle drie gebruiken dezelfde rapportstructuur en dezelfde grafiekrenderer als
het scherm. De PDF tekent de pagina's van ``ReportLayout`` rechtstreeks
vectorieel; het Excel-dashboard plakt exact dezelfde grafieken als
hoogwaardige afbeeldingen, met daarachter de onderliggende cijfers.
"""
from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path

from PySide6.QtCore import QMarginsF
from PySide6.QtGui import QPageLayout, QPageSize, QPainter, QPdfWriter

from emt_charts import render_trend_image
from emt_report import datum_tekst, page_orientation
from emt_report_layout import ReportLayout
from emt_trends import available_dimensions, build_series, ranked_events, series_totals


class ExportError(RuntimeError):
    """Een fout die de gebruiker in gewone woorden te zien krijgt."""


# ------------------------------------------------------------------- PDF

def export_pdf(model: dict, path, orientation: str = "") -> Path:
    """Schrijf het rapport als PDF, met dezelfde pagina's als het voorbeeld."""
    path = Path(path)
    layout = ReportLayout(model, orientation or page_orientation(model))
    writer = QPdfWriter(str(path))
    writer.setResolution(72)  # één eenheid is één punt, net als in de opmaak
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setPageOrientation(
        QPageLayout.Orientation.Landscape if layout.orientation == "liggend"
        else QPageLayout.Orientation.Portrait
    )
    writer.setPageMargins(QMarginsF(0, 0, 0, 0), QPageLayout.Unit.Point)
    writer.setTitle(str(model.get("titel", "") or "Trendsrapport"))

    painter = QPainter()
    if not painter.begin(writer):
        raise ExportError("Het PDF-bestand kon niet worden geopend om naar te schrijven.")
    try:
        for index in range(layout.page_count()):
            if index:
                writer.newPage()
            layout.paint_page(painter, index)
    finally:
        painter.end()
    return path


# ------------------------------------------------------------------- PNG

def export_chart_png(series: dict, path, width: int = 1200, height: int = 460,
                     scale: float = 2.0) -> Path:
    """Eén grafiek als afbeelding, met dezelfde stijl als op het scherm."""
    path = Path(path)
    image = render_trend_image(series, width, height, dark=False, scale=scale)
    if not image.save(str(path)):
        raise ExportError("De afbeelding kon niet worden opgeslagen.")
    return path


# ----------------------------------------------------------------- Excel

def _chart_blocks(model: dict) -> list[dict]:
    blokken = []
    for sectie in model.get("sections", []):
        for blok in sectie["blocks"]:
            if blok["kind"] == "chart":
                blokken.append(blok)
    return blokken


def _sheet_name(naam: str, gebruikt: set) -> str:
    """Excel accepteert 31 tekens en geen : \\ / ? * [ ]."""
    schoon = "".join(teken for teken in str(naam) if teken not in set(':\\/?*[]'))[:31] or "Blad"
    kandidaat, nummer = schoon, 2
    while kandidaat.lower() in gebruikt:
        achtervoegsel = f" {nummer}"
        kandidaat = schoon[:31 - len(achtervoegsel)] + achtervoegsel
        nummer += 1
    gebruikt.add(kandidaat.lower())
    return kandidaat


def export_excel(model: dict, path) -> Path:
    """Een werkmap met een visueel dashboard én de cijfers erachter."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    path = Path(path)
    workbook = Workbook()
    gebruikt = set()
    dashboard = workbook.active
    dashboard.title = _sheet_name("Samenvatting", gebruikt)
    dashboard.sheet_view.showGridLines = False

    accent = Font(name="Segoe UI", size=18, bold=True, color="1F1147")
    zacht = Font(name="Segoe UI", size=9, color="637187")
    kop = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    waarde_font = Font(name="Segoe UI", size=16, bold=True, color="17233A")
    vulling = PatternFill("solid", fgColor="EEF3F9")
    kop_vulling = PatternFill("solid", fgColor="6C2CFF")
    rand = Border(*[Side(style="thin", color="D8E0EB")] * 4)

    for kolom in range(1, 13):
        dashboard.column_dimensions[get_column_letter(kolom)].width = 15

    dashboard["B2"] = str(model.get("titel", "Trendsrapport"))
    dashboard["B2"].font = accent
    dashboard["B3"] = str(model.get("subtitel", "") or model.get("periode", ""))
    dashboard["B3"].font = zacht
    regel = 4
    for naam, waarde in model.get("filters", []):
        dashboard.cell(row=regel, column=2, value=f"{naam}: {waarde}").font = zacht
        regel += 1
    dashboard.cell(row=regel, column=2,
                   value=f"Gegenereerd op {datum_tekst(date.today())}").font = zacht

    # Kerncijfers als tegels.
    kpis = model.get("kpis", {})
    tegels = [
        ("Evenementen", kpis.get("evenementen", 0), "0"),
        ("Aanmeldingen", kpis.get("aangemeld", 0), "#,##0"),
        ("Aanwezigen", kpis.get("aanwezig", 0), "#,##0"),
        ("No-shows", kpis.get("noshows", 0), "#,##0"),
        ("Gem. opkomst", (kpis.get("opkomst_percentage", 0) or 0) / 100, "0,0%"),
        ("No-showpercentage", (kpis.get("noshow_percentage", 0) or 0) / 100, "0,0%"),
    ]
    top = regel + 2
    for index, (label, waarde, opmaak) in enumerate(tegels):
        kolom = 2 + (index % 3) * 3
        rij = top + (index // 3) * 3
        cel_label = dashboard.cell(row=rij, column=kolom, value=label)
        cel_label.font = zacht
        cel_label.fill = vulling
        cel_waarde = dashboard.cell(row=rij + 1, column=kolom, value=waarde)
        cel_waarde.font = waarde_font
        cel_waarde.number_format = opmaak
        cel_waarde.fill = vulling
        for verschuiving in (0, 1, 2):
            for hoogte in (0, 1):
                cel = dashboard.cell(row=rij + hoogte, column=kolom + verschuiving)
                cel.fill = vulling
                cel.border = rand
        dashboard.merge_cells(start_row=rij, start_column=kolom,
                              end_row=rij, end_column=kolom + 2)
        dashboard.merge_cells(start_row=rij + 1, start_column=kolom,
                              end_row=rij + 1, end_column=kolom + 2)

    # De grafieken: exact dezelfde renders als in Trends en de PDF.
    tijdelijk = Path(tempfile.mkdtemp(prefix="eventhub-rapport-"))
    beeldrij = top + ((len(tegels) + 2) // 3) * 3 + 1
    try:
        from openpyxl.drawing.image import Image as ExcelImage

        for index, blok in enumerate(_chart_blocks(model)):
            bestand = tijdelijk / f"grafiek{index}.png"
            render_trend_image(blok.get("series", {}), 760, 300, dark=False, scale=2.0).save(str(bestand))
            titel = dashboard.cell(row=beeldrij, column=2, value=str(blok.get("title", "")))
            titel.font = Font(name="Segoe UI", size=11, bold=True, color="17233A")
            afbeelding = ExcelImage(str(bestand))
            afbeelding.width, afbeelding.height = 760, 300
            dashboard.add_image(afbeelding, f"B{beeldrij + 1}")
            beeldrij += 18
    except ImportError:
        dashboard.cell(row=beeldrij, column=2,
                       value="Grafieken vereisen de module Pillow.").font = zacht

    # ------------------------------------------------------------ datatabbladen
    summaries = model.get("summaries", []) or []

    def blad(naam, kolommen, rijen, formats=None):
        if not rijen:
            return
        sheet = workbook.create_sheet(_sheet_name(naam, gebruikt))
        sheet.append(list(kolommen))
        for cel in sheet[1]:
            cel.font = kop
            cel.fill = kop_vulling
            cel.alignment = Alignment(vertical="center")
        for rij in rijen:
            sheet.append(list(rij))
        for kolom_index, opmaak in (formats or {}).items():
            for rij in range(2, sheet.max_row + 1):
                sheet.cell(row=rij, column=kolom_index).number_format = opmaak
        for kolom_index in range(1, len(kolommen) + 1):
            langste = max(
                [len(str(kolommen[kolom_index - 1]))]
                + [len(str(rij[kolom_index - 1])) for rij in rijen]
            )
            sheet.column_dimensions[get_column_letter(kolom_index)].width = min(46, langste + 4)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(kolommen))}{sheet.max_row}"

    from emt_models import parse_date

    evenementen = []
    for item in ranked_events(summaries):
        stats = item.get("statistiek", {})
        evenementen.append([
            str(item.get("name", "")),
            parse_date(item.get("date", "")),
            str(item.get("event_type", "")),
            str(item.get("location", "") or item.get("place", "")),
            int(stats.get("aangemeld", 0) or 0),
            int(stats.get("aanwezig", 0) or 0),
            int(stats.get("noshows", 0) or 0),
            (item.get("turnout", 0.0) or 0.0) / 100,
        ])
    blad("Evenementen",
         ["Evenement", "Datum", "Soort", "Locatie", "Aanmeldingen", "Aanwezig",
          "No-shows", "Opkomst"],
         evenementen,
         {2: "dd-mm-yyyy", 5: "#,##0", 6: "#,##0", 7: "#,##0", 8: "0,0%"})

    from emt_report import chart_period
    period = chart_period(model.get("selectie", {}))
    opkomst_reeks = build_series(summaries, "opkomst_percentage", "", period)
    noshow_reeks = build_series(summaries, "noshow_percentage", "", period)
    aanmeld_reeks = build_series(summaries, "aangemeld", "", period)
    aanwezig_reeks = build_series(summaries, "aanwezig", "", period)
    per_periode = []
    noshow_per_periode = {punt["label"]: punt for punt in noshow_reeks.get("points", [])}
    aanmeld_per_periode = {punt["label"]: punt for punt in aanmeld_reeks.get("points", [])}
    aanwezig_per_periode = {punt["label"]: punt for punt in aanwezig_reeks.get("points", [])}
    for punt in opkomst_reeks.get("points", []):
        label = punt["label"]
        groep = (opkomst_reeks.get("groups") or ["Totaal"])[0]
        per_periode.append([
            label,
            int(aanmeld_per_periode.get(label, {}).get("values", {}).get(groep, 0)),
            int(aanwezig_per_periode.get(label, {}).get("values", {}).get(groep, 0)),
            punt["values"].get(groep, 0.0) / 100,
            noshow_per_periode.get(label, {}).get("values", {}).get(groep, 0.0) / 100,
        ])
    blad("Opkomst",
         ["Periode", "Aanmeldingen", "Aanwezig", "Opkomst", "No-showpercentage"],
         per_periode, {2: "#,##0", 3: "#,##0", 4: "0,0%", 5: "0,0%"})

    verdelingen = set(available_dimensions(summaries))
    for dimensie, bladnaam in (("Opleidingsniveau", "Opleiding"), ("Profiel", "Profielen"),
                               ("Geslacht", "Demografie"), ("Leeftijdsgroep", "Leeftijd")):
        if model.get("preset") == "huidige_analyse" and dimensie != model.get("selectie", {}).get("analysis_dimension"):
            continue
        if dimensie not in verdelingen:
            continue
        totalen = sorted(series_totals(build_series(summaries, "aangemeld", dimensie, "year")),
                         key=lambda item: item[1], reverse=True)
        opkomst = dict(series_totals(build_series(summaries, "opkomst_percentage", dimensie, "year")))
        alles = sum(waarde for _naam, waarde in totalen) or 1
        selected_groups = model.get("selectie", {}).get("group_selections", {}).get(dimensie)
        if selected_groups is not None:
            totalen = [(name, value) for name, value in totalen if name in selected_groups]
        blad(bladnaam, [dimensie, "Aanmeldingen", "Aandeel", "Opkomst"],
             [[naam, int(waarde), waarde / alles, opkomst.get(naam, 0.0) / 100]
              for naam, waarde in totalen],
             {2: "#,##0", 3: "0,0%", 4: "0,0%"})

    try:
        workbook.save(str(path))
    except PermissionError as exc:
        raise ExportError(
            "Het bestand is nog geopend in Excel. Sluit het en probeer het opnieuw."
        ) from exc
    return path
