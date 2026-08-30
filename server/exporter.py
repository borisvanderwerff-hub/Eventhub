"""Bezoekerslijst export for EventHub Server.

Self-contained (no dependency on EventHub Desktop's export templates,
consistent with the rest of this package): writes a plain, readable
.xlsx with one row per participant, using openpyxl directly.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="071A33")
_HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

COLUMNS = [
    ("Voornaam", "voornaam", 18),
    ("Tussenvoegsel", "tussenvoegsel", 14),
    ("Achternaam", "achternaam", 20),
    ("Geboortedatum", "geboortedatum", 15),
    ("Geboorteplaats", "geboorteplaats", 20),
    ("Geslacht", "geslacht", 12),
    ("Opleidingsniveau", "opleidingsniveau", 16),
    ("Profiel", "profiel", 18),
    ("Telefoonnummer", "telefoonnummer", 16),
    ("E-mail", "email", 24),
    ("Introducee", "introducee", 12),
    ("Aanwezig", "aanwezig", 12),
    ("Ingecheckt om", "checkin_time", 20),
    ("Ingecheckt door", "checkin_by", 18),
]

_STATUS_LABELS = {
    "present": "Aanwezig",
    "not_checked_in": "Nog niet ingecheckt",
}


def export_participants_workbook(participants: list[dict], event: dict | None, output_path: str | Path,
                                 selected_fields: list[str] | None = None) -> int:
    """Write the participant list to an .xlsx file. Returns the row count written."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bezoekerslijst"

    event = event or {}
    sheet.append([event.get("name", "Evenement")])
    sheet.append([f"{event.get('date', '')} — {event.get('location', '')}"])
    sheet["A1"].font = Font(name="Segoe UI", size=13, bold=True)
    sheet["A2"].font = Font(name="Segoe UI", size=10)
    sheet.append([])

    header_row = sheet.max_row + 1
    columns = [column for column in COLUMNS if selected_fields is None or column[1] in selected_fields]
    if not columns:
        raise ValueError("Selecteer minimaal één kolom.")
    for col_index, (label, _, _) in enumerate(columns, start=1):
        cell = sheet.cell(header_row, col_index, label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    row_count = 0
    for participant in participants:
        row_values = []
        for _, field, _ in columns:
            value = participant.get(field, "")
            if field == "introducee":
                value = "Ja" if value else "Nee"
            elif field == "aanwezig":
                value = "Ja" if participant.get("checkin_time") else "Nee"
            row_values.append(value if value is not None else "")
        sheet.append(row_values)
        row_count += 1

    for col_index, (_, _, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(col_index)].width = width
    sheet.freeze_panes = sheet.cell(header_row + 1, 1).coordinate

    workbook.save(output_path)
    return row_count
