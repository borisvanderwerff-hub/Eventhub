import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from copy import deepcopy
import tempfile
from pathlib import Path
import unittest
from PySide6.QtWidgets import QApplication
from bezoekerslijst_app import TrendPanel
from emt_trends import build_series, select_series_groups
from emt_charts import group_colour
from emt_report import default_config, build_report_model
from emt_report_layout import ReportLayout
from emt_report_export import export_excel


def rows():
    return [{"id": str(i), "name": f"Editie {i}", "date": f"01-0{i + 1}-2025",
             "statistiek": {"aangemeld": 100, "aanwezig": 70, "noshows": 20, "afgemeld": 10,
                "verdeling": {"Opleidingsniveau": {
                    "MBO 3": {"aangemeld": 20, "aanwezig": 14, "noshows": 4, "afgemeld": 2},
                    "MBO 4": {"aangemeld": 30, "aanwezig": 21, "noshows": 6, "afgemeld": 3},
                    "HBO": {"aangemeld": 50, "aanwezig": 35, "noshows": 10, "afgemeld": 5}}}}}
            for i in range(3)]


class FocusedTrendsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_selection_preserves_denominators_and_input(self):
        series = build_series(rows(), "aangemeld", "Opleidingsniveau", "month", percentage_of_total=True)
        original = deepcopy(series)
        result = select_series_groups(series, {"Opleidingsniveau": ["MBO 3", "MBO 4"]})
        self.assertEqual(set(result["groups"]), {"MBO 3", "MBO 4"})
        self.assertEqual(result["points"][0]["values"]["MBO 3"], 20)
        self.assertEqual(series, original)
        self.assertEqual(select_series_groups(series, {"Opleidingsniveau": []})["groups"], [])

    def test_time_unit_survives_refresh_and_statistics_toggle(self):
        panel = TrendPanel(rows)
        panel.refresh()
        panel.time_choice.setCurrentIndex(panel.time_choice.findData("quarter"))
        panel.refresh()
        self.assertEqual(panel.chart.series["period"], "quarter")
        panel.display_choice.setCurrentIndex(panel.display_choice.findData("statistieken"))
        self.assertEqual(panel.chart.series["period"], "total")
        panel.display_choice.setCurrentIndex(0)
        self.assertEqual(panel.chart.series["period"], "quarter")
        panel.close()

    def test_current_analysis_export_is_focused(self):
        config = default_config("huidige_analyse")
        config["selectie"].update(analysis_metric="aangemeld", analysis_dimension="Opleidingsniveau",
            group_selections={"Opleidingsniveau": ["MBO 3"]}, periode="quarter")
        model = build_report_model(rows(), config)
        self.assertEqual(len(model["sections"]), 1)
        charts = [b for s in model["sections"] for b in s["blocks"] if b["kind"] == "chart"]
        self.assertEqual(charts[0]["series"]["groups"], ["MBO 3"])
        self.assertEqual(model["kpis"]["aangemeld"], 300)
        with tempfile.TemporaryDirectory() as directory:
            from openpyxl import load_workbook
            path = export_excel(model, Path(directory) / "test.xlsx")
            book = load_workbook(path)
            self.assertEqual(book["Opkomst"].max_row, 2)
            self.assertEqual(book["Opleiding"].max_row, 2)
            self.assertEqual(book["Opleiding"].cell(2, 3).value, 0.2)
            book.close()

    def test_long_table_text_increases_row_height_and_paginates(self):
        block = {"kind": "table", "columns": ["Opleiding", "Aantal"],
                 "rows": [["Uitgebreide opleiding met een lange naam " * 5, "30"] for _ in range(70)], "weights": [2, 1]}
        model = {"sections": [{"title": "Opleiding", "blocks": [block]}], "vormgeving": {"voorblad": False}}
        layout = ReportLayout(model)
        self.assertGreater(layout.page_count(), 2)
        copied = [b for p in layout.pages for *_, b in p["items"] if b["kind"] == "table"]
        self.assertEqual(sum(len(b["rows"]) for b in copied), 70)
        for page in layout.pages:
            for _, y, _, height, _ in page["items"]:
                self.assertLessEqual(y + height, layout.content_bottom + 0.1)

    def test_colours_are_stable(self):
        before = group_colour("MBO 3").name()
        group_colour("HBO")
        self.assertEqual(group_colour("MBO 3").name(), before)
        self.assertNotEqual(before, group_colour("MBO 4").name())

    def test_excel_keeps_more_than_six_selected_charts(self):
        model = build_report_model(rows(), default_config("volledig"))
        count = sum(block["kind"] == "chart" for section in model["sections"] for block in section["blocks"])
        self.assertGreater(count, 6)
        with tempfile.TemporaryDirectory() as directory:
            from openpyxl import load_workbook
            path = export_excel(model, Path(directory) / "all.xlsx")
            book = load_workbook(path)
            self.assertEqual(len(book["Samenvatting"]._images), count)
            book.close()
