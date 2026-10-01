from __future__ import annotations

from copy import copy
from datetime import date, datetime
from pathlib import Path
import re
import unicodedata
import zipfile
from xml.sax.saxutils import escape as xml_escape

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
import xlrd


from emt_attendance import (
    ABSENT_VALUES, AANWEZIG, AFGEMELD, AFWEZIG, ATTENDANCE_LABELS,
    ATTENDANCE_STATUSES, CALLBACK_DONE_STATUSES, CALLBACK_STATUSES,
    CANCELLED_VALUES, DUBBELE_INSCHRIJVING, INSCHRIJVING, ONBEKEND,
    OVERGESLAGEN, PRESENT_VALUES, UNKNOWN_VALUES, attendance_counts,
    attendance_map, attendance_status, attendance_value, callback_is_done,
    callback_status, infer_attendance_from_text, infer_presence_from_text,
    is_cancelled, is_introducee, is_no_show, is_present, is_present_in_scope,
    normalize, record_events, rename_attendance_event, set_attendance,
    set_present, strongest_status, text, turnout_percentage, _STATUS_GEWICHT,
)

STRING_FIELDS = [
    "Evenement", "Identifier", "GastVan", "Voornaam", "Tussenvoegsel", "Achternaam",
    "Email", "Telefoonnummer", "Geboortedatum", "Geboorteplaats", "Geslacht",
    "Opleiding", "Profiel", "Type", "Aanwezigheid", "Gebruik",
]


FIELD_ALIASES = {
    "Evenement": [
        "Evenement", "Event", "Naam evenement", "Evenementnaam", "Eventnaam",
        "Titel evenement", "Evenement titel", "Event title", "Eventtitel",
        "Naam event", "Activiteit", "Activiteitsnaam", "Bijeenkomst",
    ],
    "Identifier": ["Identifier", "ID", "Kandidaatnummer", "Registratienummer", "Registratie ID", "Registratie-ID"],
    "GastVan": ["Gast van", "Gastvan", "Guest of", "Introducé van", "Introducee van"],
    "Voornaam": ["Voornaam", "First name", "Firstname"],
    "Tussenvoegsel": ["Tussenvoegsel", "Prefix"],
    "Achternaam": ["Achternaam", "Last name", "Lastname"],
    "Email": ["E-mail", "Email", "E-mailadres"],
    "Telefoonnummer": ["Telefoonnummer", "Telefoon", "Mobiel", "Mobiele nummer", "Phone"],
    "Geboortedatum": ["Geboortedatum", "Date of birth", "Birth date"],
    "Geboorteplaats": [
        "Geboorteplaats", "Geboorte plaats", "Plaats van geboorte",
        "Plaats geboorte", "Geboortestad", "Geb. plaats", "Geb plaats",
        "Geboorteplaats volgens identiteitsbewijs", "Geboorteplaats ID",
        "Place of birth", "Birth place", "Birthplace", "Birth city",
        "City of birth", "Birth location",
    ],
    "Geslacht": ["Geslacht", "Gender"],
    "Opleiding": ["Opleiding", "Opleidingsniveau", "Niveau", "Education level"],
    "Profiel": [
        "Opleidingsrichting", "Profiel", "Studierichting",
        "Opleidingsprofiel", "Richting", "Study profile",
    ],
    "Type": ["Type"],
    "Aanwezigheid": [
        "Aanwezigheid", "Aanwezig", "Present", "Presence", "Attendance",
        "Opgekomen", "Show", "No show", "Noshow",
    ],
    "Gebruik": ["Gebruik"],
}


# De vier standen die een aanmelding kan hebben. Tot bestandsversie 11 was dit
# een ja/nee, waarin 'nee' zowel 'was er niet' als 'we weten het niet' betekende.
# Die samenklap was de bron van stille fouten: een import zonder ingevulde
# kolom en een livesessie zonder scans maakten allebei van iedereen een
# no-show. Met vier standen valt er niets meer te raden.
# Van welke aanmeldpagina een deelnemer kwam. Dezelfde dag staat soms onder
# meerdere namen online om verschillende doelgroepen te trekken; na het
# bundelen is dit het enige wat daarvan overblijft, en juist die informatie is
# de reden dat die aparte pagina's bestaan.

# Waarom een deelnemer nergens meer meetelt. Leeg betekent gewoon meetellen.
# Dit is bewust een reden en geen ja/nee: zonder reden is later niet meer te
# zien waarom iemand uit de lijst is verdwenen.


_STATUS_GEWICHT = {ONBEKEND: 0, AFGEMELD: 1, AFWEZIG: 2, AANWEZIG: 3}


EVENT_NAME_DATE_SUFFIX = re.compile(
    r"\s*(?:\(\d{2}[_-]\d{2}[_-]\d{4}\)|\(\d{2}-\d{2}-'\d{2}\)|—\s*\d{2}-\d{2}-\d{4})\s*$"
)


def event_base_name(name: str) -> str:
    """De evenementnaam zonder de datum die EventHub erachter zet."""
    return EVENT_NAME_DATE_SUFFIX.sub("", str(name or "").strip()).strip()


def _home_for_orphan_attendance(orphan: str, linked_events: list[str]) -> str:
    """Bij welk evenement hoort aanwezigheid die onder een losse naam staat?"""
    if len(linked_events) == 1:
        # De bezoeker hangt maar aan een evenement; ergens anders kan het niet horen.
        return linked_events[0]
    wanted = normalize(event_base_name(orphan))
    matches = [name for name in linked_events if normalize(event_base_name(name)) == wanted]
    return matches[0] if len(matches) == 1 else ""


def clear_absence_for_events(records: list[dict], event_names) -> int:
    """Zet 'niet gekomen' terug naar 'onbekend' bij deze evenementen.

    Bedoeld voor dossiers van voor bestandsversie 12. Daarin was er maar een
    ja/nee, en alles wat geen ja was werd een nee. Bij een evenement dat nog
    moet plaatsvinden is dat aantoonbaar onjuist: er kan nog niemand zijn
    weggebleven. Wie al als aanwezig geregistreerd staat blijft ongemoeid.
    """
    wanted = {normalize(name) for name in event_names if str(name or "").strip()}
    if not wanted:
        return 0
    changed = 0
    for record in records:
        presence = attendance_map(record)
        aangepast = False
        for name, status in list(presence.items()):
            if status == AFWEZIG and normalize(name) in wanted:
                presence[name] = ONBEKEND
                aangepast = True
        if aangepast:
            record["Aanwezig"] = presence
            changed += 1
    return changed


def repair_orphan_attendance(records: list[dict], event_names) -> int:
    """Breng aanwezigheid onder een naam die geen evenement is alsnog thuis.

    Lijsten die voor deze versie zijn ingelezen legden de aanwezigheid vast
    onder de evenementnaam uit de bron. Die naam mist de datum die EventHub
    erbij zet, en soms heet de aanmeldlijst zelfs heel anders dan het
    evenement. Zolang zo'n sleutel bij geen enkel evenement hoort, telt die
    aanwezigheid nergens mee: iedereen staat als no-show.

    Verhuizen gebeurt alleen als er een ondubbelzinnig tehuis is, en een
    registratie die er al staat wordt nooit teruggedraaid.
    """
    known = {normalize(name) for name in event_names if str(name or "").strip()}
    repaired = 0
    for record in records:
        presence = attendance_map(record)
        orphans = [name for name in presence if normalize(name) not in known]
        if not orphans:
            continue
        linked = [name for name in record_events(record) if normalize(name) in known]
        changed = False
        for orphan in orphans:
            home = _home_for_orphan_attendance(orphan, linked)
            if not home:
                continue
            gevonden = presence.pop(orphan)
            key = next((name for name in presence if normalize(name) == normalize(home)), home)
            bestaand = presence.get(key, ONBEKEND)
            # Wat er al staat blijft staan, behalve wanneer het nog onbekend is
            # of wanneer de gevonden stand aanwezigheid bevestigt.
            if bestaand == ONBEKEND or gevonden == AANWEZIG:
                presence[key] = gevonden
            else:
                presence[key] = bestaand
            changed = True
        if changed:
            record["Aanwezig"] = presence
            repaired += 1
    return repaired


def detach_event_from_records(records: list[dict], event_name: str) -> list[dict]:
    """Remove one event association while preserving multi-event visitors."""
    wanted = normalize(event_name)
    retained = []
    for record in records:
        linked_names = record_events(record)
        if not any(normalize(name) == wanted for name in linked_names):
            retained.append(record)
            continue
        remaining_names = [name for name in linked_names if normalize(name) != wanted]
        if remaining_names:
            record["Evenement"] = "; ".join(remaining_names)
            # Laat de aanwezigheid van het losgekoppelde evenement niet achter.
            record["Aanwezig"] = {
                name: value for name, value in attendance_map(record).items()
                if normalize(name) != wanted
            }
            retained.append(record)
    return retained


def has_status_in_scope(record: dict, scope_events, status: str) -> bool:
    """Heeft deze bezoeker die stand bij een van de evenementen in beeld?"""
    names = list(scope_events) if scope_events else record_events(record)
    return any(attendance_status(record, name) == status for name in names)


def _name_words(name: str) -> list[str]:
    return [word for word in event_base_name(name).split() if word]


def common_event_name(names) -> str:
    """Het deel dat alle namen delen; de voorgestelde naam voor de bundel.

    Meeloopdag Marine Catering, Administratie en Logistiek delen Meeloopdag
    Marine. Delen ze niets, dan valt hij terug op de eerste naam.
    """
    woordenlijsten = [_name_words(name) for name in names if str(name or "").strip()]
    if not woordenlijsten:
        return ""
    gemeen: list[str] = []
    for positie in range(min(len(woorden) for woorden in woordenlijsten)):
        kandidaat = woordenlijsten[0][positie]
        if all(normalize(woorden[positie]) == normalize(kandidaat) for woorden in woordenlijsten):
            gemeen.append(kandidaat)
        else:
            break
    return " ".join(gemeen) or event_base_name(next(iter(names)))


def distinctive_labels(names) -> dict:
    """Wat elke naam uniek maakt: Catering, Administratie, Logistiek.

    Het gedeelde deel valt weg, zodat het label kort blijft in de statistieken.
    Blijft er niets over, dan is de hele naam het label - dan is er niets
    onderscheidends en is de volledige naam het eerlijkste antwoord.
    """
    namen = [str(name or "").strip() for name in names if str(name or "").strip()]
    gedeeld = len(_name_words(common_event_name(namen))) if len(namen) > 1 else 0
    labels = {}
    for naam in namen:
        rest = " ".join(_name_words(naam)[gedeeld:]).strip()
        labels[naam] = rest or event_base_name(naam)
    return labels


def is_skipped(record: dict) -> bool:
    """Telt en toont deze regel nog mee?"""
    return bool(str(record.get(OVERGESLAGEN, "") or "").strip())


def skip_reason(record: dict) -> str:
    return str(record.get(OVERGESLAGEN, "") or "").strip()


def counting_records(records) -> list:
    """Alles wat meetelt; overgeslagen regels vallen hier af."""
    return [record for record in records or [] if not is_skipped(record)]


def richest_record(records) -> dict | None:
    """Welke regel het meeste weet.

    Bij een dubbele inschrijving blijft die staan: hij heeft de meeste kans
    de juiste gegevens te bevatten, en aanwezigheid weegt daarbij het zwaarst.
    """
    kandidaten = [record for record in records or [] if record is not None]
    if not kandidaten:
        return None

    def gewicht(record):
        gevuld = sum(1 for field in STRING_FIELDS if text(record.get(field)))
        return (AANWEZIG in attendance_map(record).values(), gevuld)

    return max(kandidaten, key=gewicht)


def absorb_duplicate(keeper: dict, duplicate: dict, reason: str = DUBBELE_INSCHRIJVING) -> None:
    """Laat de blijvende regel overnemen wat de dubbele weet, en sla die over.

    Zonder dat overnemen verdwijnt met de dubbele regel ook zijn aanwezigheid:
    bij een van de dubbele deelnemers stond de aanwezigheid juist op de regel
    die zou worden verborgen, en dan telde hij ineens als no-show.

    De regel blijft in het dossier staan. Weghalen zou hem bij de volgende
    import gewoon terugbrengen; overslaan houdt stand.
    """
    if keeper is None or duplicate is None or keeper is duplicate:
        return
    for field in STRING_FIELDS:
        if not text(keeper.get(field)) and text(duplicate.get(field)):
            keeper[field] = text(duplicate.get(field))
    for event_name, status in attendance_map(duplicate).items():
        set_attendance(keeper, event_name,
                       strongest_status(attendance_status(keeper, event_name), status))
    for label in registrations(duplicate):
        add_registration(keeper, label)
    duplicate[OVERGESLAGEN] = str(reason or DUBBELE_INSCHRIJVING).strip() or DUBBELE_INSCHRIJVING


def include_again(record: dict) -> bool:
    """Laat een overgeslagen regel weer meetellen."""
    if not is_skipped(record):
        return False
    record[OVERGESLAGEN] = ""
    return True


def registrations(record: dict) -> list[str]:
    """Via welke aanmeldpagina's deze deelnemer binnenkwam."""
    raw = str(record.get(INSCHRIJVING, "") or "")
    return [part.strip() for part in raw.split(";") if part.strip()]


def add_registration(record: dict, label: str) -> bool:
    """Noteer de aanmeldpagina, zonder dubbele vermeldingen."""
    label = str(label or "").strip()
    if not label:
        return False
    bestaand = registrations(record)
    if any(normalize(name) == normalize(label) for name in bestaand):
        return False
    record[INSCHRIJVING] = "; ".join(bestaand + [label])
    return True


def merge_event_into(records: list[dict], source_name: str, target_name: str,
                     label: str = "") -> int:
    """Verhuis de koppeling en de aanwezigheid van het ene evenement naar het andere.

    Anders dan relink_to_event blijven andere evenementen van dezelfde bezoeker
    ongemoeid: alleen het samengevoegde evenement verandert van naam. Stond
    iemand al bij het doelevenement, dan wint de meest uitgesproken stand.
    """
    source_name = str(source_name or "").strip()
    target_name = str(target_name or "").strip()
    if not source_name or not target_name:
        return 0
    verplaatst = 0
    for record in records:
        namen = record_events(record)
        if not any(normalize(name) == normalize(source_name) for name in namen):
            continue
        presence = attendance_map(record)
        uit_bron = presence.pop(
            next((name for name in presence if normalize(name) == normalize(source_name)), ""),
            ONBEKEND,
        )
        doel_sleutel = next(
            (name for name in presence if normalize(name) == normalize(target_name)), target_name
        )
        presence[doel_sleutel] = strongest_status(presence.get(doel_sleutel, ONBEKEND), uit_bron)
        record["Aanwezig"] = presence

        vernieuwd: list[str] = []
        for name in namen:
            nieuwe = target_name if normalize(name) == normalize(source_name) else name
            if not any(normalize(nieuwe) == normalize(bestaand) for bestaand in vernieuwd):
                vernieuwd.append(nieuwe)
        record["Evenement"] = "; ".join(vernieuwd)
        add_registration(record, label or source_name)
        verplaatst += 1
    return verplaatst


def merge_duplicate_registrations(records: list[dict], event_name: str) -> tuple[list[dict], int]:
    """Dezelfde persoon uit twee aanmeldlijsten wordt één deelnemer.

    Bij losse evenementen zag je zo iemand helemaal niet; na het bundelen staat
    hij er twee keer in. Zijn beide inschrijvingen blijven bewaard, zodat
    zichtbaar blijft dat hij zich twee keer heeft aangemeld.
    """
    wanted = normalize(event_name)
    gezien: dict = {}
    behouden: list[dict] = []
    samengevoegd = 0
    for record in records:
        hoort_erbij = any(normalize(name) == wanted for name in record_events(record))
        key = duplicate_key(record) if hoort_erbij else None
        eerste = gezien.get(key) if key else None
        if eerste is None:
            if key:
                gezien[key] = record
            behouden.append(record)
            continue
        # Alles wat de dubbele regel extra weet, gaat mee naar de eerste.
        _enrich_record(eerste, record)
        for naam, status in attendance_map(record).items():
            eerste["Aanwezig"] = attendance_map(eerste)
            set_attendance(eerste, naam, strongest_status(attendance_status(eerste, naam), status))
        for label in registrations(record):
            add_registration(eerste, label)
        samengevoegd += 1
    return behouden, samengevoegd


def matches_event_filter(record: dict, selected_events) -> bool:
    if not selected_events:
        return True
    return bool(set(record_events(record)) & set(selected_events))


def visitor_type(record: dict) -> str:
    return "Introducé" if is_introducee(record) else "Reguliere bezoeker"


def registration_lookup(records) -> dict[str, dict]:
    return {
        normalize(record.get("Identifier", "")): record
        for record in records
        if normalize(record.get("Identifier", ""))
    }


def primary_visitor_name(record: dict, lookup: dict[str, dict]) -> str:
    if not is_introducee(record):
        return ""
    primary = lookup.get(normalize(record.get("GastVan", "")))
    if not primary:
        return "Onbekende hoofdbezoeker"
    return " ".join(filter(None, [
        text(primary.get("Voornaam")),
        text(primary.get("Tussenvoegsel")),
        text(primary.get("Achternaam")),
    ]))


def date_text(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).strftime("%d-%m-%Y")
        except (ValueError, OverflowError, TypeError):
            pass
    return text(value)


def initials_text(value) -> str:
    """Maak nette voorletters van een volledige voornaam."""
    parts = re.findall(r"[^\W\d_]+", text(value), flags=re.UNICODE)
    return "".join(f"{part[0].upper()}." for part in parts if part)


def participant_template_context(event: dict | None, profile: dict | None = None) -> dict[str, str]:
    event = event or {}
    profile = profile or {}
    fivewh = event.get("fivewh", {}) if isinstance(event.get("fivewh"), dict) else {}
    start = text(fivewh.get("event_time", ""))
    departure = text(fivewh.get("teardown_time", ""))
    if start and departure:
        times = f"{start} / vertrek {departure}"
    else:
        times = start or departure
    return {
        "event_name": text(event.get("name", "")),
        "event_date": date_text(event.get("date", "")),
        "event_times": times,
        "event_location": " — ".join(filter(None, [
            text(event.get("place", "")), text(event.get("location", "")),
        ])),
        "contact_name": text(event.get("external_contact", "")) or text(profile.get("name", "")),
        "contact_phone": text(event.get("external_contact_reachability", "")) or text(profile.get("phone", "")),
    }


def phone_text(value) -> str:
    result = text(value)
    if result.endswith(".0") and result[:-2].isdigit():
        result = result[:-2]
    if len(result) == 9 and result.isdigit() and not result.startswith("0"):
        result = f"0{result}"
    return result


def duplicate_key(record: dict) -> str:
    if record.get("Identifier"):
        return f"id:{normalize(record['Identifier'])}"
    if record.get("Email"):
        return f"mail:{normalize(record['Email'])}"
    digits = re.sub(r"\D", "", record.get("Telefoonnummer", ""))
    if len(digits) >= 8:
        return f"tel:{digits}"
    return "naam:" + normalize("|".join([
        record.get("Voornaam", ""), record.get("Tussenvoegsel", ""),
        record.get("Achternaam", ""), record.get("Geboortedatum", ""),
    ]))


HEADER_PATTERNS = {
    "Evenement": ("evenement", "eventnaam", "eventtitel", "eventtitle", "naamevent", "activiteitsnaam"),
    "Geboorteplaats": (
        "geboorteplaats", "plaatsvangeboorte", "plaatsgeboorte", "geboortestad",
        "gebplaats", "placeofbirth", "birthplace", "birthcity", "cityofbirth", "birthlocation",
    ),
}

def _first_column(value) -> int:
    from emt_registration_import import _first_column as _implementation
    return _implementation(value)


def build_map(headers: dict[str, list[int] | int]) -> dict[str, int]:
    from emt_registration_import import build_map as _implementation
    return _implementation(headers)


def _looks_like_date_value(value) -> bool:
    from emt_registration_import import _looks_like_date_value as _implementation
    return _implementation(value)


def _duplicate_birthplace_column(headers, rows, header_row_index, mapping) -> int:
    from emt_registration_import import _duplicate_birthplace_column as _implementation
    return _implementation(headers, rows, header_row_index, mapping)


def _xlsx_sheets(file_path: Path):
    from emt_registration_import import _xlsx_sheets as _implementation
    return _implementation(file_path)


def _xls_sheets(file_path: Path):
    from emt_registration_import import _xls_sheets as _implementation
    return _implementation(file_path)


def read_registration_file(file_path: str | Path):
    from emt_registration_import import read_registration_file as _implementation
    return _implementation(file_path)


def relink_to_event(record: dict, event_name: str) -> None:
    from emt_registration_import import relink_to_event as _implementation
    return _implementation(record, event_name)


def _merge_attendance(target: dict, incoming: dict, overwrite: bool=False, conflicts: list | None=None) -> bool:
    from emt_registration_import import _merge_attendance as _implementation
    return _implementation(target, incoming, overwrite, conflicts)


def _enrich_record(target: dict, incoming: dict, overwrite: bool=False, conflicts: list | None=None) -> bool:
    from emt_registration_import import _enrich_record as _implementation
    return _implementation(target, incoming, overwrite, conflicts)


def import_registration_files(file_paths, existing_records=None, target_event: str='', overwrite_attendance: bool=False):
    from emt_registration_import import import_registration_files as _implementation
    return _implementation(file_paths, existing_records, target_event, overwrite_attendance)


def apply_attendance_conflicts(conflicts) -> int:
    from emt_registration_import import apply_attendance_conflicts as _implementation
    return _implementation(conflicts)


def _style_sheet(worksheet, widths, wrap_columns=()):
    from emt_excel_export import _style_sheet as _implementation
    return _implementation(worksheet, widths, wrap_columns)


def export_participant_template(records, event: dict, profile: dict, template_path: str | Path, output_path: str | Path):
    from emt_excel_export import export_participant_template as _implementation
    return _implementation(records, event, profile, template_path, output_path)


def export_workbook(records, output_path: str | Path, reference_records=None):
    from emt_excel_export import export_workbook as _implementation
    return _implementation(records, output_path, reference_records)


KOP_VULLING = PatternFill("solid", fgColor="071A33")


def _titel(sheet, regel: int, tekst: str, grootte: int=13):
    from emt_excel_export import _titel as _implementation
    return _implementation(sheet, regel, tekst, grootte)


def _koprij(sheet, regel: int, koppen, breedtes):
    from emt_excel_export import _koprij as _implementation
    return _implementation(sheet, regel, koppen, breedtes)


def _write_summary_sheet(workbook, summary: dict) -> None:
    from emt_excel_export import _write_summary_sheet as _implementation
    return _implementation(workbook, summary)


def export_statistics_workbook(dimensions, output_path: str | Path, summary: dict | None=None, crosstab: dict | None=None):
    from emt_excel_export import export_statistics_workbook as _implementation
    return _implementation(dimensions, output_path, summary, crosstab)


def _heatmap_fill(fraction: float) -> PatternFill:
    from emt_excel_export import _heatmap_fill as _implementation
    return _implementation(fraction)


def _write_crosstab_sheet(workbook, crosstab: dict, header_fill) -> None:
    from emt_excel_export import _write_crosstab_sheet as _implementation
    return _implementation(workbook, crosstab, header_fill)
