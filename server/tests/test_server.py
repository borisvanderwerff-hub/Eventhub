"""Automated tests for EventHub Server.

Run with:  python -m unittest server.tests.test_server -v
(from the repository root, e.g. EventHub_Source_v2.8.1/)

Uses only the standard library's unittest (no pytest dependency) plus
openpyxl/Flask, which are already required by the project.
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
import threading
import unittest
from pathlib import Path

# Isolate all file-system state for the test run before importing server.*.
_TEST_DATA_ROOT = Path(tempfile.mkdtemp(prefix="eventhub_server_tests_"))
os.environ["XDG_DATA_HOME"] = str(_TEST_DATA_ROOT)

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

import openpyxl  # noqa: E402

from server import importer  # noqa: E402
from server.database import close_all, connect  # noqa: E402
from server.services import client_service, participant_service, session_service, statistics_service  # noqa: E402
from server.web.app import create_app  # noqa: E402
from server import qrgen  # noqa: E402
from server.exporter import export_participants_workbook  # noqa: E402


def _make_test_workbook(path: Path, rows: list[list]):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Voornaam", "Achternaam", "Geboortedatum", "Opleiding", "Profiel", "Geslacht", "Gast van"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


class SessionValidationTests(unittest.TestCase):
    def test_rejects_empty_fields(self):
        with self.assertRaises(session_service.SessionValidationError):
            session_service.create_session("", "24-08-2026", "Ergens")
        with self.assertRaises(session_service.SessionValidationError):
            session_service.create_session("Naam", "", "Ergens")
        with self.assertRaises(session_service.SessionValidationError):
            session_service.create_session("Naam", "24-08-2026", "")

    def test_creates_session_with_uuid(self):
        result = session_service.create_session("Open Dag", "24-08-2026", "Hoofdgebouw")
        self.assertTrue(result.id)
        self.assertTrue(result.db_path.is_file())
        connection = connect(result.db_path)
        session = session_service.get_session(connection)
        self.assertEqual(session["name"], "Open Dag")
        self.assertEqual(session["status"], "active")

    def test_creates_a_four_digit_event_code(self):
        result = session_service.create_session("Open Dag", "24-08-2026", "Hoofdgebouw")
        self.assertRegex(result.event_code, r"^\d{4}$")
        connection = connect(result.db_path)
        session = session_service.get_session(connection)
        self.assertEqual(session["event_code"], result.event_code)

    def test_verify_event_code(self):
        result = session_service.create_session("Codetest", "24-08-2026", "Locatie")
        connection = connect(result.db_path)
        self.assertTrue(session_service.verify_event_code(connection, result.event_code))
        self.assertFalse(session_service.verify_event_code(connection, "0000"))
        self.assertFalse(session_service.verify_event_code(connection, ""))


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.session = session_service.create_session("Importtest", "24-08-2026", "Locatie")
        self.connection = connect(self.session.db_path)
        self.workbook_path = _TEST_DATA_ROOT / "import_test.xlsx"

    def test_reuses_core_column_detection_and_marks_not_checked_in(self):
        _make_test_workbook(self.workbook_path, [
            ["Jan", "Jansen", "01-01-1995", "MBO", "ICT", "Man", ""],
            ["Marie", "de Vries", "02-02-1998", "HBO", "Zorg", "Vrouw", ""],
            ["Tom", "Tomassen", "03-03-2000", "WO", "Techniek", "Man", "Jan Jansen"],
        ])
        preview = importer.preview_import(self.workbook_path)
        self.assertEqual(preview["row_count"], 3)
        self.assertIn("Opleiding", preview["detected_columns"])
        self.assertEqual(preview["introducees"], 1)

        rows = importer.records_to_participants(preview["records"], self.session.id)
        self.assertTrue(all(row["attendance_status"] == "not_checked_in" for row in rows))
        count = participant_service.add_participants(self.connection, self.session.id, rows)
        self.assertEqual(count, 3)

        participants = participant_service.list_participants(self.connection, self.session.id)
        self.assertEqual(len(participants), 3)
        self.assertTrue(all(p["attendance_status"] == "not_checked_in" for p in participants))
        introducees = [p for p in participants if p["introducee"]]
        self.assertEqual(len(introducees), 1)

    def test_detects_duplicates_against_existing_participants(self):
        _make_test_workbook(self.workbook_path, [["Jan", "Jansen", "01-01-1995", "MBO", "ICT", "Man", ""]])
        preview = importer.preview_import(self.workbook_path)
        rows = importer.records_to_participants(preview["records"], self.session.id)
        participant_service.add_participants(self.connection, self.session.id, rows)

        existing = participant_service.list_participants(self.connection, self.session.id)
        preview_again = importer.preview_import(self.workbook_path, existing_participants=existing)
        self.assertEqual(preview_again["duplicate_count"], 1)

    def test_ignores_source_present_column_per_spec_section_7(self):
        # Even if bezoekerslijst_core detects a Present/Aanwezig column and
        # pre-fills record['Aanwezig'], the Server must not use it: importing
        # only registers someone, it never marks them present.
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Voornaam", "Achternaam", "Present"])
        sheet.append(["Jan", "Jansen", "yes"])
        workbook.save(self.workbook_path)

        preview = importer.preview_import(self.workbook_path)
        self.assertEqual(preview["presence_detected"], 1)  # core did detect it...
        rows = importer.records_to_participants(preview["records"], self.session.id)
        self.assertEqual(rows[0]["attendance_status"], "not_checked_in")  # ...but server ignores it


class CheckinTests(unittest.TestCase):
    def setUp(self):
        self.session = session_service.create_session("Checkintest", "24-08-2026", "Locatie")
        self.connection = connect(self.session.db_path)
        rows = importer.records_to_participants(
            [{"Voornaam": "Jan", "Achternaam": "Jansen", "Geboortedatum": "01-01-1995"}],
            self.session.id,
        )
        participant_service.add_participants(self.connection, self.session.id, rows)
        self.participant_id = rows[0]["id"]

    def test_checkin_then_idempotent_second_checkin(self):
        result = participant_service.check_in(self.connection, self.participant_id)
        self.assertFalse(result["already_checked_in"])
        self.assertEqual(result["participant"]["attendance_status"], "present")

        result_again = participant_service.check_in(self.connection, self.participant_id)
        self.assertTrue(result_again["already_checked_in"])
        # No duplicate audit entries beyond the single 'checked_in' for two checkin calls.
        count = self.connection.execute(
            "SELECT COUNT(*) AS n FROM audit_log WHERE action = 'checked_in' AND participant_id = ?",
            (self.participant_id,),
        ).fetchone()["n"]
        self.assertEqual(count, 1)

    def test_checkout_then_idempotent_second_checkout(self):
        participant_service.check_in(self.connection, self.participant_id)
        result = participant_service.check_out(self.connection, self.participant_id)
        self.assertFalse(result["already_not_present"])
        self.assertEqual(result["participant"]["attendance_status"], "not_checked_in")

        result_again = participant_service.check_out(self.connection, self.participant_id)
        self.assertTrue(result_again["already_not_present"])

    def test_concurrent_checkins_only_count_once(self):
        errors = []

        def do_checkin():
            try:
                participant_service.check_in(self.connection, self.participant_id)
            except Exception as exc:  # pragma: no cover - defensive
                errors.append(exc)

        threads = [threading.Thread(target=do_checkin) for _ in range(8)]
        for t in threads:
            t.start()
        for t in threads:
            t.join(timeout=10)

        self.assertFalse(errors, f"Onverwachte fouten tijdens gelijktijdige check-ins: {errors}")
        count = self.connection.execute(
            "SELECT COUNT(*) AS n FROM audit_log WHERE action = 'checked_in' AND participant_id = ?",
            (self.participant_id,),
        ).fetchone()["n"]
        self.assertEqual(count, 1, "Een deelnemer mag maar één keer als ingecheckt gelogd worden.")
        participant = participant_service.get_participant(self.connection, self.participant_id)
        self.assertEqual(participant["attendance_status"], "present")


class StatisticsTests(unittest.TestCase):
    def setUp(self):
        self.session = session_service.create_session("Statstest", "24-08-2026", "Locatie")
        self.connection = connect(self.session.db_path)
        rows = importer.records_to_participants([
            {"Voornaam": "A", "Achternaam": "A", "Opleiding": "MBO", "Geslacht": "Man"},
            {"Voornaam": "B", "Achternaam": "B", "Opleiding": "MBO", "Geslacht": "Vrouw"},
            {"Voornaam": "C", "Achternaam": "C", "Opleiding": "HBO", "Geslacht": "Man"},
        ], self.session.id)
        participant_service.add_participants(self.connection, self.session.id, rows)
        participant_service.check_in(self.connection, rows[0]["id"])
        participant_service.check_in(self.connection, rows[1]["id"])
        self.rows = rows

    def test_overview_counts_and_turnout(self):
        overview = statistics_service.overview(self.connection, self.session.id)
        self.assertEqual(overview["registered"], 3)
        self.assertEqual(overview["present"], 2)
        self.assertEqual(overview["expected"], 1)
        self.assertAlmostEqual(overview["turnout_percentage"], 66.7, places=1)

    def test_breakdown_by_education_reports_per_group_turnout(self):
        breakdown = statistics_service.breakdown_by_dimension(self.connection, self.session.id, "education")
        by_label = {item["label"]: item for item in breakdown}
        self.assertEqual(by_label["MBO"]["registered"], 2)
        self.assertEqual(by_label["MBO"]["present"], 2)
        self.assertEqual(by_label["MBO"]["turnout_percentage"], 100.0)
        self.assertEqual(by_label["HBO"]["registered"], 1)
        self.assertEqual(by_label["HBO"]["present"], 0)

    def test_client_checkin_counts(self):
        client = client_service.register_client(self.connection, self.session.id, "Balie 1", "Browser", "127.0.0.1", "checkin")
        participant_service.check_in(self.connection, self.rows[2]["id"], client_id=client["id"], client_name="Balie 1")
        counts = statistics_service.client_checkin_counts(self.connection, self.session.id)
        self.assertEqual(counts[0]["client_name"], "Balie 1")
        self.assertEqual(counts[0]["checkin_count"], 1)


class PersistenceTests(unittest.TestCase):
    def test_session_and_participants_survive_reconnect(self):
        session = session_service.create_session("Persistentietest", "24-08-2026", "Locatie")
        connection = connect(session.db_path)
        rows = importer.records_to_participants(
            [{"Voornaam": "Jan", "Achternaam": "Jansen"}], session.id
        )
        participant_service.add_participants(connection, session.id, rows)
        participant_service.check_in(connection, rows[0]["id"])

        # Simulate an application restart: close every connection, then
        # reopen the same database file from scratch.
        close_all()
        reconnected = connect(session.db_path)
        resumed_session = session_service.get_session(reconnected)
        self.assertEqual(resumed_session["id"], session.id)
        participants = participant_service.list_participants(reconnected, session.id)
        self.assertEqual(len(participants), 1)
        self.assertEqual(participants[0]["attendance_status"], "present")

    def test_resume_lists_session_in_registry(self):
        session = session_service.create_session("Hervattest", "24-08-2026", "Locatie")
        recent = session_service.list_recent_sessions()
        self.assertTrue(any(item["id"] == session.id for item in recent))
        db_path = session_service.resume_session(session.id)
        self.assertEqual(db_path, session.db_path)


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.session = session_service.create_session("APItest", "24-08-2026", "Locatie")
        self.app = create_app(self.session.db_path, self.session.id)
        self.client = self.app.test_client()

    def _register(self, role="admin"):
        response = self.client.post("/api/clients/register", json={
            "client_name": "Testbalie", "client_type": "Browser", "role": role,
            "event_code": self.session.event_code,
        })
        self.assertEqual(response.status_code, 201)
        registered = response.get_json()
        self.assertEqual(registered["role"], "viewer")
        if role != "viewer":
            connection = connect(self.session.db_path)
            registered = client_service.set_role(
                connection, self.session.id, registered["id"], role, "Desktopbeheerder"
            )
        return registered

    def test_health_and_server_info(self):
        self.assertEqual(self.client.get("/api/health").get_json()["status"], "ok")
        info = self.client.get("/api/server").get_json()
        self.assertEqual(info["product"], "EventHub Server")
        self.assertEqual(info["event"]["id"], self.session.id)
        self.assertNotIn("event_code", info["event"])  # never leak the code over the API

    def test_registration_requires_correct_event_code(self):
        response = self.client.post("/api/clients/register", json={
            "client_name": "Testbalie", "client_type": "Browser", "role": "checkin",
            "event_code": "0000",
        })
        self.assertEqual(response.status_code, 401)

        response = self.client.post("/api/clients/register", json={
            "client_name": "Testbalie", "client_type": "Browser", "role": "checkin",
        })
        self.assertEqual(response.status_code, 401)

        response = self.client.post("/api/clients/register", json={
            "client_name": "Testbalie", "client_type": "Browser", "role": "checkin",
            "event_code": self.session.event_code,
        })
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.get_json()["role"], "viewer")

    def test_client_can_disconnect_logout(self):
        client = self._register(role="checkin")
        response = self.client.delete(f"/api/clients/{client['id']}", headers={"X-Client-Id": client["id"]})
        self.assertEqual(response.status_code, 200)
        clients = self.client.get("/api/clients").get_json()
        self.assertFalse(any(c["id"] == client["id"] for c in clients))

    def test_mutating_endpoints_require_client_registration(self):
        response = self.client.post("/api/import/preview", data={}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 401)

    def test_participant_list_requires_client_registration(self):
        self.assertEqual(self.client.get("/api/participants").status_code, 401)

    def test_checkin_role_enforcement(self):
        rows = importer.records_to_participants([{"Voornaam": "Jan", "Achternaam": "Jansen"}], self.session.id)
        connection = connect(self.session.db_path)
        participant_service.add_participants(connection, self.session.id, rows)
        viewer = self._register(role="viewer")
        response = self.client.post(f"/api/participants/{rows[0]['id']}/checkin",
                                     headers={"X-Client-Id": viewer["id"]})
        self.assertEqual(response.status_code, 403)

    def test_checkin_requires_registered_client(self):
        rows = importer.records_to_participants([{"Voornaam": "Jan", "Achternaam": "Jansen"}], self.session.id)
        connection = connect(self.session.db_path)
        participant_service.add_participants(connection, self.session.id, rows)
        response = self.client.post(f"/api/participants/{rows[0]['id']}/checkin")
        self.assertEqual(response.status_code, 401)

    def test_full_import_and_checkin_flow_via_http(self):
        admin = self._register(role="admin")
        workbook_path = _TEST_DATA_ROOT / "api_test.xlsx"
        _make_test_workbook(workbook_path, [["Jan", "Jansen", "01-01-1995", "MBO", "ICT", "Man", ""]])
        with open(workbook_path, "rb") as file_handle:
            response = self.client.post(
                "/api/import/preview",
                data={"file": (file_handle, "bezoekers.xlsx")},
                headers={"X-Client-Id": admin["id"]},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        preview = response.get_json()
        self.assertEqual(preview["file_name"], "bezoekers.xlsx")

        response = self.client.post("/api/import/commit", json={"import_token": preview["import_token"]},
                                     headers={"X-Client-Id": admin["id"]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["imported"], 1)

        participants = self.client.get(
            "/api/participants", headers={"X-Client-Id": admin["id"]}
        ).get_json()
        self.assertEqual(len(participants), 1)
        participant_id = participants[0]["id"]

        response = self.client.post(f"/api/participants/{participant_id}/checkin",
                                     headers={"X-Client-Id": admin["id"]})
        self.assertEqual(response.status_code, 201)

        overview = self.client.get("/api/statistics/overview").get_json()
        self.assertEqual(overview["present"], 1)
        self.assertEqual(overview["registered"], 1)


class QrCodeTests(unittest.TestCase):
    def test_generates_valid_square_matrix(self):
        matrix = qrgen.generate_matrix("http://192.168.10.15:8080", ec_level="M")
        self.assertTrue(all(len(row) == len(matrix) for row in matrix))
        self.assertIn(len(matrix), (21, 25, 29, 33))  # version 1-4 sizes

    def test_rejects_text_too_long_for_supported_versions(self):
        with self.assertRaises(ValueError):
            qrgen.generate_matrix("x" * 200, ec_level="M")

    def test_svg_output_contains_expected_number_of_dark_modules(self):
        matrix = qrgen.generate_matrix("http://eventhub.local:8080", ec_level="M")
        svg = qrgen.matrix_to_svg(matrix)
        dark_count = sum(1 for row in matrix for cell in row if cell)
        self.assertEqual(svg.count("<rect x="), dark_count)


class ExportTests(unittest.TestCase):
    def test_exports_participants_with_status_and_introducee_labels(self):
        session = session_service.create_session("Exporttest", "24-08-2026", "Locatie")
        connection = connect(session.db_path)
        rows = importer.records_to_participants([
            {"Voornaam": "Jan", "Achternaam": "Jansen", "Opleiding": "MBO"},
            {"Voornaam": "Marie", "Achternaam": "de Vries", "Opleiding": "HBO", "GastVan": "Jan Jansen"},
        ], session.id)
        participant_service.add_participants(connection, session.id, rows)
        participant_service.check_in(connection, rows[0]["id"])

        participants = participant_service.list_participants(connection, session.id)
        event = session_service.get_session(connection)
        output_path = _TEST_DATA_ROOT / "export_test.xlsx"
        count = export_participants_workbook(participants, event, output_path)
        self.assertEqual(count, 2)

        workbook = openpyxl.load_workbook(output_path)
        sheet = workbook.active
        rows_out = list(sheet.iter_rows(values_only=True))
        header_row = next(row for row in rows_out if row and row[0] == "Voornaam")
        header_index = rows_out.index(header_row)
        data_rows = {row[0]: row for row in rows_out[header_index + 1:]}
        status_index = header_row.index("Status")
        introducee_index = header_row.index("Introducee")
        self.assertEqual(data_rows["Jan"][status_index], "Aanwezig")
        self.assertEqual(data_rows["Marie"][status_index], "Nog niet ingecheckt")
        self.assertEqual(data_rows["Marie"][introducee_index], "Ja")


class _CleanupMixin:
    @classmethod
    def tearDownClass(cls):
        close_all()


def load_tests(loader, tests, pattern):
    suite = unittest.TestSuite()
    for test_case in (SessionValidationTests, ImportTests, CheckinTests, StatisticsTests,
                       PersistenceTests, ApiTests, QrCodeTests, ExportTests):
        suite.addTests(loader.loadTestsFromTestCase(test_case))
    return suite


if __name__ == "__main__":
    try:
        unittest.main(verbosity=2)
    finally:
        close_all()
        shutil.rmtree(_TEST_DATA_ROOT, ignore_errors=True)
