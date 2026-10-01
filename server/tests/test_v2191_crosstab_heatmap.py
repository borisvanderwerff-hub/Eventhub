"""De kruistabel opleidingsniveau x profiel als heatmap, op het scherm en in de export.

Het raster is in de praktijk dertien niveaus bij achtentwintig profielen en
voor ruim tachtig procent leeg. Als tabel lees je vooral nullen. De staart van
de profielen wordt daarom gebundeld, lege cellen blijven leeg en de kleur zegt
hoe groot het aantal is.
"""
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
from server.tests.desktop_source import desktop_source
sys.path.insert(0, str(ROOT))

from bezoekerslijst_core import export_statistics_workbook
from emt_education import collapse_columns, crosstab

APP_SOURCE = desktop_source(ROOT)


def deelnemer(niveau, profiel):
    return {"Opleiding": niveau, "Profiel": profiel}


class CollapseColumnsTests(unittest.TestCase):
    def _brede_tabel(self):
        """Twee grote profielen en een staart van kleintjes."""
        records = [deelnemer("MBO 4", "Techniek") for _ in range(10)]
        records += [deelnemer("HAVO", "Economie") for _ in range(6)]
        records += [deelnemer("VWO", f"Klein profiel {index}") for index in range(6)]
        return crosstab(records)

    def test_de_staart_wordt_een_kolom(self):
        data = collapse_columns(self._brede_tabel(), limit=2)

        self.assertEqual(data["columns"], ["Techniek", "Economie", "Overig (6)"])

    def test_er_gaat_geen_deelnemer_verloren(self):
        origineel = self._brede_tabel()
        data = collapse_columns(origineel, limit=2)

        self.assertEqual(sum(data["column_totals"].values()), origineel["total"])
        self.assertEqual(sum(data["counts"].values()), origineel["total"])

    def test_de_gebundelde_namen_blijven_bekend(self):
        """Anders kun je niet nagaan wat er onder Overig zit."""
        data = collapse_columns(self._brede_tabel(), limit=2)

        self.assertEqual(len(data["gebundeld"]), 6)
        self.assertIn("Klein profiel 3", data["gebundeld"])

    def test_de_grootste_profielen_blijven_apart(self):
        data = collapse_columns(self._brede_tabel(), limit=2)

        self.assertEqual(data["column_totals"]["Techniek"], 10)
        self.assertEqual(data["column_totals"]["Economie"], 6)
        self.assertEqual(data["column_totals"]["Overig (6)"], 6)

    def test_de_rijen_blijven_ongemoeid(self):
        origineel = self._brede_tabel()
        data = collapse_columns(origineel, limit=2)

        self.assertEqual(data["rows"], origineel["rows"])
        self.assertEqual(data["row_totals"], origineel["row_totals"])

    def test_een_smalle_tabel_blijft_zoals_hij_is(self):
        """Overig (1) is alleen een omweg naar hetzelfde getal."""
        data = crosstab([deelnemer("MBO 4", "Techniek"), deelnemer("HAVO", "Economie")])

        self.assertEqual(collapse_columns(data, limit=8), data)

    def test_bundelen_kan_worden_uitgezet(self):
        origineel = self._brede_tabel()

        self.assertEqual(collapse_columns(origineel, limit=0), origineel)


class CrosstabSheetTests(unittest.TestCase):
    def setUp(self):
        self.folder = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.folder, ignore_errors=True)
        records = [deelnemer("MBO 4", "Techniek") for _ in range(8)]
        records += [deelnemer("HAVO", "Economie") for _ in range(4)]
        self.data = crosstab(records)
        self.path = self.folder / "statistieken.xlsx"
        # (categorie, aangemeld, aanwezig, niet gekomen, afgemeld)
        export_statistics_workbook(
            [("Opleidingsniveau", [("MBO 4", 8, 8, 0, 0), ("HAVO", 4, 4, 0, 0)])],
            self.path,
            crosstab=self.data,
        )
        self.sheet = openpyxl.load_workbook(self.path)["Niveau x profiel"]

    def _cel(self, kop_rij, niveau, profiel):
        kolom = [c.value for c in self.sheet[kop_rij]].index(profiel) + 1
        for row in self.sheet.iter_rows(min_row=kop_rij + 1):
            if row[0].value == niveau:
                return row[kolom - 1]
        raise AssertionError(f"{niveau} niet gevonden")

    def test_een_lege_cel_blijft_leeg(self):
        """Een nul in tachtig procent van het raster is alleen maar ruis."""
        self.assertIsNone(self._cel(3, "MBO 4", "Economie").value)

    def test_een_gevulde_cel_krijgt_kleur(self):
        cel = self._cel(3, "MBO 4", "Techniek")

        self.assertEqual(cel.value, 8)
        self.assertEqual(cel.fill.patternType, "solid")

    def test_een_lege_cel_krijgt_geen_kleur(self):
        self.assertIn(self._cel(3, "MBO 4", "Economie").fill.patternType, (None, "none"))

    def test_ook_lege_cellen_hebben_subtiele_rasterlijnen(self):
        cel = self._cel(3, "MBO 4", "Economie")
        self.assertEqual(cel.border.left.style, "thin")
        self.assertEqual(cel.border.left.color.rgb, "00D9DCE3")

    def test_er_staat_ook_een_blok_met_aandelen(self):
        koppen = [self.sheet.cell(row, 1).value for row in range(1, self.sheet.max_row + 1)]

        self.assertIn("Aantallen", koppen)
        self.assertIn("Aandeel per niveau", koppen)

    def test_het_aandeel_staat_als_percentage(self):
        aandeel_rij = [self.sheet.cell(row, 1).value for row in range(1, self.sheet.max_row + 1)]
        kop = aandeel_rij.index("Aandeel per niveau") + 1
        cel = self._cel(kop, "MBO 4", "Techniek")

        self.assertEqual(cel.value, 1.0)
        self.assertEqual(cel.number_format, "0%")

    def test_aandeelblok_toont_aantal_in_plaats_van_nietszeggend_honderd_procent(self):
        eerste_kolom = [self.sheet.cell(row, 1).value for row in range(1, self.sheet.max_row + 1)]
        kop = eerste_kolom.index("Aandeel per niveau") + 1
        koppen = [cell.value for cell in self.sheet[kop]]
        self.assertEqual(koppen[-1], "Aantal")
        mbo_rij = next(row for row in self.sheet.iter_rows(min_row=kop + 1) if row[0].value == "MBO 4")
        self.assertEqual(mbo_rij[-1].value, 8)

    def test_de_totalen_blijven_staan(self):
        totaal_rij = next(
            row for row in self.sheet.iter_rows(min_row=4) if row[0].value == "Totaal"
        )

        self.assertEqual(totaal_rij[-1].value, 12)


class CompactExportTests(unittest.TestCase):
    """De grafiek stond op een vaste regel, ver onder een korte tabel."""

    def setUp(self):
        self.folder = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.folder, ignore_errors=True)

    def _blad(self, aantal_categorieen):
        path = self.folder / f"export-{aantal_categorieen}.xlsx"
        rijen = [(f"Groep {index}", index + 1, index, 1, 0) for index in range(aantal_categorieen)]
        export_statistics_workbook([("Opleidingsniveau", rijen)], path)
        return openpyxl.load_workbook(path)["Opleidingsniveau"]

    def _grafiekrij(self, aantal_categorieen):
        return self._blad(aantal_categorieen)._charts[0].anchor._from.row + 1

    def test_er_is_nog_maar_een_grafiek_per_blad(self):
        """Drie van vijftien centimeter op tien centimeter afstand overlapten."""
        self.assertEqual(len(self._blad(6)._charts), 1)

    def test_de_grafiek_volgt_de_tabel(self):
        """Drie categorieen: rijen 4 tot 6, totaal op 7, grafiek op 9."""
        self.assertEqual(self._grafiekrij(3), 9)

    def test_een_langere_tabel_schuift_de_grafiek_op(self):
        self.assertEqual(self._grafiekrij(12), 18)

    def test_er_zit_geen_veld_lege_rijen_meer_tussen(self):
        """Voorheen stond hij altijd op rij 22, ook bij drie categorieen."""
        self.assertLess(self._grafiekrij(3), 22)


class TheScreenShowsTheHeatmapTests(unittest.TestCase):
    def test_de_heatmap_bestaat(self):
        self.assertIn("class CrosstabHeatmap(QWidget):", APP_SOURCE)

    def test_een_lege_cel_krijgt_geen_inkt(self):
        start = APP_SOURCE.index("class CrosstabHeatmap(QWidget):")
        widget = APP_SOURCE[start:APP_SOURCE.index("class TrendChart(QWidget):")]

        self.assertIn("if not count:", widget)
        self.assertIn("continue", widget)

    def test_de_tabel_blijft_bereikbaar(self):
        self.assertIn("self.crosstab_stack.addWidget(self.crosstab_heatmap)", APP_SOURCE)
        self.assertIn("self.crosstab_stack.addWidget(self.crosstab_table)", APP_SOURCE)

    def test_de_keuzes_worden_onthouden(self):
        self.assertIn('self.settings.setValue("crosstab_view"', APP_SOURCE)
        self.assertIn('self.settings.setValue("crosstab_value"', APP_SOURCE)
        # Alles tonen geldt sinds v2198 voor het hele tabblad, niet alleen de
        # kruistabel, en heet daarom statistics_show_all.
        self.assertIn('self.settings.setValue("statistics_show_all"', APP_SOURCE)

    def test_de_export_toont_dezelfde_kruistabel_als_het_scherm(self):
        self.assertIn("crosstab=self._crosstab_data(crosstab_records),", APP_SOURCE)

    def test_export_laat_de_deelnemersgroepen_expliciet_kiezen(self):
        self.assertIn("def _crosstab_export_scope_dialog", APP_SOURCE)
        self.assertIn("for status in (AANWEZIG, AFGEMELD, AFWEZIG, ONBEKEND)", APP_SOURCE)
        self.assertIn("Introducees meenemen", APP_SOURCE)
        self.assertIn("Alleen de kruistabel", APP_SOURCE)

    def test_de_hele_export_kan_tot_een_status_worden_beperkt(self):
        for label in ("Alle groepen", "Alleen aanwezig", "Alleen afwezig", "Alleen afgemeld"):
            self.assertIn(label, APP_SOURCE)
        start = APP_SOURCE.index("def export_statistics(self)")
        handler = APP_SOURCE[start:start + 7000]
        self.assertIn('if export_status != "all":', handler)
        self.assertIn("self._scoped_status(record) == export_status", handler)

    def test_dezelfde_keuze_wordt_op_de_hele_export_toegepast(self):
        start = APP_SOURCE.index("def export_statistics(self)")
        handler = APP_SOURCE[start:start + 6500]
        filtering = "record for record in event_records if self._scoped_status(record) in crosstab_statuses"
        self.assertIn(filtering, handler)
        self.assertIn("crosstab=self._crosstab_data(crosstab_records),", handler)


class TheExportIsNamedAfterTheEventTests(unittest.TestCase):
    def test_de_bestandsnaam_bevat_het_evenement(self):
        start = APP_SOURCE.index("def export_statistics(self)")
        handler = APP_SOURCE[start:start + 7000]

        self.assertIn("event_base_name(event.get(\"name\", \"\"))", handler)
        self.assertIn('f"EventHub-statistieken - {event_label}', handler)


if __name__ == "__main__":
    unittest.main()
