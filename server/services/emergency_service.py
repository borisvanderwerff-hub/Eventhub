from __future__ import annotations

import sqlite3
import uuid
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..database import now_iso, transaction
from ..events_hub import hub
from .session_service import get_session, SessionValidationError


DEFAULT_INSTRUCTION = "Volg de instructies van de eventmanager of beheerder."


def active_incident(connection: sqlite3.Connection) -> dict | None:
    row = connection.execute(
        "SELECT * FROM emergency_incident WHERE active=1 ORDER BY started_at DESC LIMIT 1"
    ).fetchone()
    return dict(row) if row else None


def state(connection: sqlite3.Connection) -> dict:
    incident = active_incident(connection)
    if not incident:
        return {"active": False, "incident_id": "", "instruction": "", "started_at": None, "started_by": ""}
    return {
        "active": True,
        "incident_id": incident["id"],
        "instruction": incident.get("instruction") or DEFAULT_INSTRUCTION,
        "started_at": incident.get("started_at"),
        "started_by": incident.get("started_by") or "",
    }


def start(connection: sqlite3.Connection, instruction: str = "", by: str = "") -> dict:
    session = get_session(connection)
    if not session:
        raise SessionValidationError("Geen actieve sessie.")
    existing = active_incident(connection)
    if existing:
        return state(connection)
    incident_id = str(uuid.uuid4())
    timestamp = now_iso()
    instruction = (instruction or DEFAULT_INSTRUCTION).strip()
    with transaction(connection) as tx:
        tx.execute(
            "INSERT INTO emergency_incident (id,event_id,active,instruction,started_at,started_by) VALUES (?,?,1,?,?,?)",
            (incident_id, session["id"], instruction, timestamp, (by or "").strip()),
        )
        tx.execute(
            "INSERT INTO emergency_status (incident_id,participant_id) "
            "SELECT ?,id FROM participant WHERE event_id=? AND attendance_status='present'",
            (incident_id, session["id"]),
        )
        tx.execute(
            "INSERT INTO audit_log (event_id,client_id,action,timestamp,new_value) VALUES (?,?, 'emergency_started',?,?)",
            (session["id"], by or None, timestamp, instruction),
        )
    result = state(connection)
    hub.publish("emergency_started", {"event_id": session["id"], **result})
    return result


def update_instruction(connection: sqlite3.Connection, instruction: str, by: str = "") -> dict:
    incident = active_incident(connection)
    if not incident:
        raise SessionValidationError("De calamiteitenmodus is niet actief.")
    instruction = (instruction or DEFAULT_INSTRUCTION).strip()
    with transaction(connection) as tx:
        tx.execute("UPDATE emergency_incident SET instruction=? WHERE id=?", (instruction, incident["id"]))
        tx.execute(
            "INSERT INTO audit_log (event_id,client_id,action,timestamp,new_value) VALUES (?,?, 'emergency_instruction_updated',?,?)",
            (incident["event_id"], by or None, now_iso(), instruction),
        )
    result = state(connection)
    hub.publish("emergency_instruction_updated", {"event_id": incident["event_id"], **result})
    return result


def end(connection: sqlite3.Connection, by: str = "") -> dict:
    incident = active_incident(connection)
    if not incident:
        return state(connection)
    timestamp = now_iso()
    with transaction(connection) as tx:
        tx.execute(
            "UPDATE emergency_incident SET active=0,ended_at=?,ended_by=? WHERE id=?",
            (timestamp, (by or "").strip(), incident["id"]),
        )
        tx.execute(
            "INSERT INTO audit_log (event_id,client_id,action,timestamp,new_value) VALUES (?,?, 'emergency_ended',?,?)",
            (incident["event_id"], by or None, timestamp, incident["id"]),
        )
    hub.publish("emergency_ended", {
        "event_id": incident["event_id"], "incident_id": incident["id"],
        "active": False, "ended_at": timestamp,
    })
    return {"active": False, "incident_id": incident["id"], "ended_at": timestamp}


def overview(connection: sqlite3.Connection, incident_id: str = "") -> dict:
    if incident_id:
        incident = connection.execute("SELECT * FROM emergency_incident WHERE id=?", (incident_id,)).fetchone()
    else:
        incident = active_incident(connection) or connection.execute(
            "SELECT * FROM emergency_incident ORDER BY started_at DESC LIMIT 1"
        ).fetchone()
    if not incident:
        return {"active": False, "participants": [], "total": 0, "safe": 0, "unaccounted": 0}
    incident = dict(incident)
    rows = connection.execute(
        "SELECT p.id,p.voornaam,p.tussenvoegsel,p.achternaam,p.geboortedatum,p.geboorteplaats,"
        "e.safe,e.assembly_point,e.team,e.marked_at,e.marked_by "
        "FROM emergency_status e JOIN participant p ON p.id=e.participant_id "
        "WHERE e.incident_id=? ORDER BY p.achternaam,p.tussenvoegsel,p.voornaam",
        (incident["id"],),
    ).fetchall()
    participants = [dict(row) for row in rows]
    safe_count = sum(bool(row["safe"]) for row in participants)
    return {
        "active": bool(incident["active"]), "incident_id": incident["id"],
        "instruction": incident.get("instruction") or DEFAULT_INSTRUCTION,
        "started_at": incident["started_at"], "started_by": incident.get("started_by") or "",
        "ended_at": incident.get("ended_at"), "ended_by": incident.get("ended_by") or "",
        "participants": participants, "total": len(participants), "safe": safe_count,
        "unaccounted": len(participants) - safe_count,
    }


def mark_safe(connection: sqlite3.Connection, participant_id: str, safe: bool, by: str = "",
              assembly_point: str = "", team: str = "") -> dict:
    incident = active_incident(connection)
    if not incident:
        raise SessionValidationError("De calamiteitenmodus is niet actief.")
    timestamp = now_iso() if safe else None
    with transaction(connection) as tx:
        changed = tx.execute(
            "UPDATE emergency_status SET safe=?,assembly_point=?,team=?,marked_at=?,marked_by=? "
            "WHERE incident_id=? AND participant_id=?",
            (1 if safe else 0, (assembly_point or "").strip(), (team or "").strip(),
             timestamp, (by or "").strip() if safe else "", incident["id"], participant_id),
        ).rowcount
        if not changed:
            raise SessionValidationError("Deze deelnemer staat niet op de actuele calamiteitenlijst.")
        tx.execute(
            "INSERT INTO audit_log (event_id,participant_id,client_id,action,timestamp,new_value) VALUES (?,?,?, ?,?,?)",
            (incident["event_id"], participant_id, by or None,
             "emergency_safe" if safe else "emergency_safe_undone", now_iso(), assembly_point or ""),
        )
    hub.publish("emergency_participant_updated", {
        "event_id": incident["event_id"], "incident_id": incident["id"],
        "participant_id": participant_id, "safe": bool(safe),
    })
    return overview(connection)


def export_workbook(connection: sqlite3.Connection, incident_id: str = "") -> BytesIO:
    data = overview(connection, incident_id)
    if not data.get("incident_id"):
        raise SessionValidationError("Geen calamiteitenregistratie beschikbaar.")
    session = get_session(connection) or {}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Calamiteitenregistratie"
    sheet.append(["Calamiteitenregistratie", session.get("name", "")])
    sheet.append(["Gestart", data.get("started_at", ""), "Door", data.get("started_by", "")])
    sheet.append(["Beëindigd", data.get("ended_at", ""), "Door", data.get("ended_by", "")])
    sheet.append(["Instructie", data.get("instruction", "")])
    sheet.append([])
    headers = ["Achternaam", "Tussenvoegsel", "Voornaam", "Geboortedatum", "Geboorteplaats",
               "Veilig gemeld", "Verzamelplaats", "Controleteam", "Tijdstip", "Medewerker"]
    sheet.append(headers)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6C2CFF")
    for person in data["participants"]:
        sheet.append([
            person.get("achternaam", ""), person.get("tussenvoegsel", ""), person.get("voornaam", ""),
            person.get("geboortedatum", ""), person.get("geboorteplaats", ""),
            "Ja" if person.get("safe") else "Nee", person.get("assembly_point", ""),
            person.get("team", ""), person.get("marked_at", ""), person.get("marked_by", ""),
        ])
    widths = [24, 16, 22, 16, 20, 16, 22, 20, 25, 22]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = f"A{header_row + 1}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
