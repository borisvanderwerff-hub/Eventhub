"""Excel registration-file reading and participant import workflows."""

from datetime import date, datetime
from pathlib import Path
import re

import xlrd
from openpyxl import load_workbook

# Imported lazily by bezoekerslijst_core's compatibility wrappers. Keeping the
# dependency here lets existing callers continue importing from the core module.
from bezoekerslijst_core import (
    FIELD_ALIASES, HEADER_PATTERNS, STRING_FIELDS, ONBEKEND, attendance_map, attendance_status,
    date_text, duplicate_key, infer_attendance_from_text, is_introducee,
    normalize, phone_text, record_events, set_attendance, text,
)


def _first_column(value) -> int:
    if isinstance(value, list):
        return value[0] if value else -1
    return int(value)

def build_map(headers: dict[str, list[int] | int]) -> dict[str, int]:
    result = {}
    for field, aliases in FIELD_ALIASES.items():
        result[field] = -1
        for alias in aliases:
            key = normalize(alias)
            if key in headers:
                result[field] = _first_column(headers[key])
                break
        if result[field] < 0:
            patterns = HEADER_PATTERNS.get(field, ())
            for header, columns in headers.items():
                if any(pattern in header for pattern in patterns):
                    result[field] = _first_column(columns)
                    break
    return result

def _looks_like_date_value(value) -> bool:
    if isinstance(value, (date, datetime)):
        return True
    if isinstance(value, (int, float)):
        return 1 <= value <= 100000
    value = text(value)
    if not value:
        return False
    patterns = (
        r"^\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}$",
        r"^\d{4}[-/.]\d{1,2}[-/.]\d{1,2}$",
    )
    return any(re.match(pattern, value) for pattern in patterns)

def _duplicate_birthplace_column(headers, rows, header_row_index, mapping) -> int:
    if mapping.get("Geboorteplaats", -1) >= 0:
        return -1
    birthdate_columns = []
    for alias in FIELD_ALIASES["Geboortedatum"]:
        columns = headers.get(normalize(alias), [])
        if not isinstance(columns, list):
            columns = [columns]
        birthdate_columns.extend(columns)
    birthdate_columns = sorted(set(birthdate_columns))
    if len(birthdate_columns) < 2:
        return -1

    first_birthdate = mapping.get("Geboortedatum", birthdate_columns[0])
    candidates = [column for column in birthdate_columns if column != first_birthdate]
    candidates.sort(key=lambda column: (column != first_birthdate + 1, column))
    for column in candidates:
        values = []
        for row in rows[header_row_index + 1:header_row_index + 41]:
            if column < len(row) and text(row[column]):
                values.append(row[column])
        if not values:
            return column
        date_values = sum(_looks_like_date_value(value) for value in values)
        if date_values / len(values) < 0.5:
            return column
    return -1

def _xlsx_sheets(file_path: Path):
    workbook = load_workbook(file_path, read_only=True, data_only=True, keep_vba=False)
    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            yield worksheet.title, rows
    finally:
        workbook.close()

def _xls_sheets(file_path: Path):
    workbook = xlrd.open_workbook(file_path)
    for worksheet in workbook.sheets():
        rows = []
        for row_index in range(worksheet.nrows):
            row = []
            for column_index in range(worksheet.ncols):
                cell = worksheet.cell(row_index, column_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = datetime(*xlrd.xldate_as_tuple(value, workbook.datemode))
                row.append(value)
            rows.append(row)
        yield worksheet.name, rows

def read_registration_file(file_path: str | Path):
    file_path = Path(file_path)
    sheets = list(_xls_sheets(file_path) if file_path.suffix.lower() == ".xls" else _xlsx_sheets(file_path))
    best = None
    for sheet_index, (sheet_name, rows) in enumerate(sheets):
        for row_index, row in enumerate(rows[:25]):
            headers = {}
            for column_index, value in enumerate(row):
                key = normalize(value)
                if key:
                    headers.setdefault(key, []).append(column_index)
            mapping = build_map(headers)
            if mapping["Voornaam"] < 0 or mapping["Achternaam"] < 0:
                continue
            repaired_birthplace_column = _duplicate_birthplace_column(headers, rows, row_index, mapping)
            if repaired_birthplace_column >= 0:
                mapping["Geboorteplaats"] = repaired_birthplace_column
            score = sum(index >= 0 for index in mapping.values())
            if best is None or score > best["score"]:
                best = {
                    "sheet_index": sheet_index,
                    "sheet_name": sheet_name,
                    "row_index": row_index,
                    "rows": rows,
                    "mapping": mapping,
                    "score": score,
                    "repaired_birthplace_header": repaired_birthplace_column >= 0,
                }
    if best is None:
        raise ValueError("Geen werkblad met de kolommen Voornaam en Achternaam gevonden.")

    records = []
    fallback_events = 0
    missing_birth_places = 0
    presence_detected = 0
    fallback_event = file_path.stem
    mapping = best["mapping"]

    for source in best["rows"][best["row_index"] + 1:]:
        def get(field):
            column = mapping[field]
            return source[column] if 0 <= column < len(source) else ""

        first_name = text(get("Voornaam"))
        last_name = text(get("Achternaam"))
        if not first_name and not last_name:
            continue
        record = {field: text(get(field)) for field in STRING_FIELDS}
        record.update({
            "Voornaam": first_name,
            "Achternaam": last_name,
            "Geboortedatum": date_text(get("Geboortedatum")),
            "Telefoonnummer": phone_text(get("Telefoonnummer")),
            "Teruggebeld": False,
            "Terugbelstatus": "Nog bellen",
            "LaatsteContact": "",
            "TerugbellenOp": "",
            "Opmerkingen": "",
            "WhatsAppStatus": "Nog te sturen",
            "WhatsAppGeopendOp": "",
            "WhatsAppVerzondenOp": "",
            "Aanwezig": {},
        })
        if not record["Evenement"]:
            record["Evenement"] = fallback_event
            fallback_events += 1
        # Pas na het vaststellen van het evenement: aanwezigheid wordt per
        # evenement vastgelegd en heeft die naam dus nodig.
        detected_presence = infer_attendance_from_text(record.get("Aanwezigheid", ""))
        if detected_presence is not None:
            for event_name in record_events(record):
                set_attendance(record, event_name, detected_presence)
            if detected_presence != ONBEKEND:
                presence_detected += 1
        if not record["Geboorteplaats"]:
            missing_birth_places += 1
        records.append(record)

    return {
        "records": records,
        "report": {
            "file_name": file_path.name,
            "sheet_name": best["sheet_name"],
            # Welke evenementnamen in dit ene bestand staan. Zonder deze
            # koppeling is na het inlezen niet meer te zien uit welk bestand
            # een evenement kwam, en dus ook niet welk Rudder-nummer erbij hoort.
            "events": sorted({name for record in records for name in record_events(record)}),
            "introducees": sum(is_introducee(record) for record in records),
            "fallback_events": fallback_events,
            "missing_birth_places": missing_birth_places,
            "presence_detected": presence_detected,
            "repaired_birthplace_header": best.get("repaired_birthplace_header", False),
        },
    }

def relink_to_event(record: dict, event_name: str) -> None:
    """Zet een ingelezen regel op het evenement waarin de lijst wordt geimporteerd.

    De bronkolom noemt het evenement zonder datum, EventHub zet die er wel bij.
    Zonder deze verhuizing blijft de aanwezigheid onder de naam uit het bestand
    staan en telt iedereen als afwezig.
    """
    uit_bestand = next(
        (status for status in attendance_map(record).values() if status != ONBEKEND),
        ONBEKEND,
    )
    record["Evenement"] = event_name
    record["Aanwezig"] = {}
    set_attendance(record, event_name, uit_bestand)

def _merge_attendance(target: dict, incoming: dict, overwrite: bool = False,
                      conflicts: list | None = None) -> bool:
    """Neem de stand uit het bestand over bij een bezoeker die er al staat.

    Waar EventHub nog niets weet, geldt het bestand. Waar allebei iets zeggen
    en dat verschilt, blijft standaard staan wat er is: een registratie die
    hier is gemaakt gooi je niet stilzwijgend weg. Zo'n botsing komt in
    ``conflicts`` terecht, zodat er alsnog om een keuze gevraagd kan worden.
    """
    changed = False
    for event_name, status in attendance_map(incoming).items():
        if status == ONBEKEND:
            continue
        huidig = attendance_status(target, event_name)
        if huidig == status:
            continue
        if huidig == ONBEKEND or overwrite:
            if set_attendance(target, event_name, status):
                changed = True
        elif conflicts is not None:
            conflicts.append((target, event_name, status))
    return changed

def _enrich_record(target: dict, incoming: dict, overwrite: bool = False,
                   conflicts: list | None = None) -> bool:
    # De aanwezigheid hoort net zo goed bij het aanvullen als de tekstvelden;
    # zonder deze regel verdwijnt die van iedereen die al in EventHub stond.
    changed = _merge_attendance(target, incoming, overwrite, conflicts)
    for field in STRING_FIELDS:
        current = text(target.get(field))
        new_value = text(incoming.get(field))
        if not new_value:
            continue
        if not current:
            target[field] = new_value
            changed = True
        elif field == "Evenement":
            current_events = {normalize(value) for value in current.split(";") if value.strip()}
            if normalize(new_value) not in current_events:
                target[field] = f"{current}; {new_value}"
                changed = True
    return changed

def import_registration_files(file_paths, existing_records=None, target_event: str = "",
                              overwrite_attendance: bool = False):
    records = list(existing_records or [])
    keyed_records = {duplicate_key(record): record for record in records}
    added = []
    duplicates = 0
    enriched = 0
    reports = []
    errors = []
    conflicts: list = []
    for file_path in file_paths:
        try:
            result = read_registration_file(file_path)
            reports.append(result["report"])
            for record in result["records"]:
                # Voor het ontdubbelen, zodat een bestaande bezoeker de
                # aanwezigheid onder de juiste evenementnaam aangevuld krijgt.
                if target_event:
                    relink_to_event(record, target_event)
                key = duplicate_key(record)
                if key in keyed_records:
                    duplicates += 1
                    if _enrich_record(keyed_records[key], record, overwrite_attendance, conflicts):
                        enriched += 1
                    continue
                keyed_records[key] = record
                added.append(record)
        except Exception as exc:
            errors.append(f"{Path(file_path).name}: {exc}")
    return {
        "records": added,
        "duplicates": duplicates,
        "enriched": enriched,
        "reports": reports,
        "errors": errors,
        # Bestaande registraties waarover het bestand iets anders zegt. Ze zijn
        # niet overschreven; de vraag daarover hoort bij de gebruiker.
        "attendance_conflicts": conflicts,
    }

def apply_attendance_conflicts(conflicts) -> int:
    """Neem alsnog over wat het bestand zei, nadat daarom is gevraagd."""
    changed = 0
    for record, event_name, status in conflicts or ():
        if set_attendance(record, event_name, status):
            changed += 1
    return changed
