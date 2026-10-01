"""Opleidingsniveau x profiel: herkenning van schrijfwijzen en de kruistabel."""
from pathlib import Path
import sys
import unittest


ROOT = Path(__file__).resolve().parents[2]
from server.tests.desktop_source import desktop_source
sys.path.insert(0, str(ROOT))

from emt_education import ONBEKEND, crosstab, education_level, level_sort_key

APP_SOURCE = desktop_source(ROOT)


def visitor(opleiding, profiel=""):
    return {"Opleiding": opleiding, "Profiel": profiel}


class LevelRecognitionTests(unittest.TestCase):
    def test_the_many_ways_to_write_mbo_4(self):
        """De aanleiding: dezelfde opleiding stond als vier losse waarden."""
        for value in ("MBO 4", "mbo-4", "MBO niveau 4", "Mbo niv. 4", "MBO4", "mbo   4"):
            with self.subTest(value=value):
                self.assertEqual(education_level(value), "MBO 4")

    def test_other_mbo_levels_stay_apart(self):
        self.assertEqual(education_level("MBO niveau 2"), "MBO 2")
        self.assertEqual(education_level("mbo3"), "MBO 3")
        self.assertEqual(education_level("MBO entree"), "MBO Entree")
        self.assertEqual(education_level("MBO"), "MBO")

    def test_vmbo_tracks_including_the_usual_abbreviations(self):
        for value, expected in (
            ("VMBO Basis", "VMBO Basis"), ("vmbo bb", "VMBO Basis"), ("VMBO-B", "VMBO Basis"),
            ("VMBO kader", "VMBO Kader"), ("vmbo kb", "VMBO Kader"), ("VMBO-K", "VMBO Kader"),
            ("VMBO GL", "VMBO Gemengd"), ("vmbo gemengde leerweg", "VMBO Gemengd"),
            ("VMBO-T", "VMBO TL"), ("vmbo theoretisch", "VMBO TL"), ("MAVO", "VMBO TL"),
            ("VMBO", "VMBO"),
        ):
            with self.subTest(value=value):
                self.assertEqual(education_level(value), expected)

    def test_higher_levels(self):
        for value, expected in (
            ("havo", "HAVO"), ("HAVO/VWO", "HAVO/VWO"), ("vwo", "VWO"),
            ("atheneum", "VWO"), ("Gymnasium", "VWO"),
            ("HBO", "HBO"), ("hbo bachelor", "HBO"), ("hogeschool", "HBO"),
            ("WO", "WO"), ("universiteit", "WO"), ("universitair", "WO"),
            ("praktijkonderwijs", "Praktijkonderwijs"),
        ):
            with self.subTest(value=value):
                self.assertEqual(education_level(value), expected)

    def test_empty_and_placeholder_values_become_unknown(self):
        for value in ("", "   ", "onbekend", "n.v.t.", "geen"):
            with self.subTest(value=value):
                self.assertEqual(education_level(value), ONBEKEND)

    def test_unrecognised_values_stay_visible(self):
        """Wegmoffelen onder Overig zou verbergen dat de bron rommelig is."""
        self.assertEqual(education_level("Zeevaartschool"), "Zeevaartschool")
        self.assertEqual(education_level("KMS"), "KMS")

    def test_levels_sort_from_low_to_high_with_unknown_last(self):
        labels = ["WO", ONBEKEND, "MBO 4", "VMBO Basis", "Zeevaartschool", "HAVO"]
        self.assertEqual(
            sorted(labels, key=level_sort_key),
            ["VMBO Basis", "MBO 4", "HAVO", "WO", "Zeevaartschool", ONBEKEND],
        )


class CrosstabTests(unittest.TestCase):
    def setUp(self):
        self.data = crosstab([
            visitor("MBO 4", "Techniek"),
            visitor("mbo-4", "techniek"),
            visitor("MBO niveau 4", "Zorg"),
            visitor("VMBO-T", "Techniek"),
            visitor("MAVO", "Zorg"),
            visitor("HAVO", "Economie"),
            visitor("", "Techniek"),
        ])

    def test_spelling_variants_land_in_one_row(self):
        self.assertEqual(self.data["row_totals"]["MBO 4"], 3)
        self.assertEqual(self.data["row_totals"]["VMBO TL"], 2)

    def test_profiles_group_case_insensitively(self):
        self.assertEqual(self.data["column_totals"]["Techniek"], 4)
        self.assertNotIn("techniek", self.data["columns"])

    def test_cells_hold_the_combination(self):
        self.assertEqual(self.data["counts"][("MBO 4", "Techniek")], 2)
        self.assertEqual(self.data["counts"][("MBO 4", "Zorg")], 1)

    def test_margins_add_up(self):
        self.assertEqual(self.data["total"], 7)
        self.assertEqual(sum(self.data["row_totals"].values()), 7)
        self.assertEqual(sum(self.data["column_totals"].values()), 7)

    def test_rows_run_from_low_to_high(self):
        self.assertEqual(self.data["rows"][:3], ["VMBO TL", "MBO 4", "HAVO"])
        self.assertEqual(self.data["rows"][-1], ONBEKEND)

    def test_merged_spellings_are_reported(self):
        """De gebruiker moet kunnen zien dat de weergave waarden samenvoegt."""
        self.assertEqual(self.data["merged"]["MBO 4"], ["MBO niveau 4", "mbo-4"])
        self.assertIn("MAVO", self.data["merged"]["VMBO TL"])

    def test_nothing_to_report_when_spelling_already_matches(self):
        data = crosstab([visitor("MBO 4", "Techniek")])
        self.assertEqual(data["merged"], {})

    def test_empty_input_is_harmless(self):
        data = crosstab([])
        self.assertEqual(data["rows"], [])
        self.assertEqual(data["total"], 0)


class IntegrationTests(unittest.TestCase):
    def test_the_existing_charts_are_untouched(self):
        """Uitdrukkelijke wens: de bestaande statistieken blijven zoals ze zijn."""
        for chart in ("education", "profile", "gender", "age"):
            self.assertIn(f'self.statistics_cards["{chart}"].chart.set_data(', APP_SOURCE)

    def test_the_crosstab_is_an_extra_card_below_them(self):
        """Onder de grafieken en over de volle breedte; welke rij precies mag schuiven."""
        kaarten = APP_SOURCE.index('statistics_grid.addWidget(self.statistics_cards["age"]')
        kruistabel = APP_SOURCE.index("statistics_grid.addWidget(crosstab_box, ")

        self.assertLess(kaarten, kruistabel)
        self.assertIn(", 0, 1, 2)", APP_SOURCE[kruistabel:kruistabel + 80])
        self.assertIn("def _update_crosstab(self, records):", APP_SOURCE)

    def test_the_crosstab_is_included_in_the_excel_export(self):
        """Op het scherm stond hij wel, in de export niet."""
        start = APP_SOURCE.index("def export_statistics(self)")
        block = APP_SOURCE[start:APP_SOURCE.index("\n    def ", start + 1)]
        self.assertIn("crosstab=self._crosstab_data(crosstab_records)", block)

    def test_the_export_button_label_is_short(self):
        """De volledige titel liep buiten de knop."""
        self.assertIn('statistics_export_button = QPushButton("Exporteren")', APP_SOURCE)
        self.assertNotIn('QPushButton("Statistieken exporteren naar Excel")', APP_SOURCE)

    def test_the_source_data_is_never_rewritten(self):
        block = APP_SOURCE[APP_SOURCE.index("def _update_crosstab(self"):]
        block = block[:block.index("\n    def ", 1)]
        self.assertNotIn('record["Opleiding"] =', block)
        self.assertNotIn('record["Profiel"] =', block)


class ExcelSheetTests(unittest.TestCase):
    """De kruistabel als eigen blad in de werkmap."""

    @classmethod
    def setUpClass(cls):
        import openpyxl
        import tempfile
        from bezoekerslijst_core import export_statistics_workbook

        records = [
            visitor("MBO 4", "Techniek"), visitor("mbo-4", "techniek"),
            visitor("MBO niveau 4", "Zorg"), visitor("HAVO", "Economie"),
        ]
        output = Path(tempfile.mkdtemp()) / "stat.xlsx"
        # (categorie, aangemeld, aanwezig, niet gekomen, afgemeld)
        export_statistics_workbook(
            [("Opleidingsniveau", [("MBO 4", 3, 3, 0, 0)])],
            output, crosstab=crosstab(records),
        )
        cls.workbook = openpyxl.load_workbook(output)

    def test_the_sheet_exists(self):
        self.assertIn("Niveau x profiel", self.workbook.sheetnames)

    def test_the_matrix_holds_the_combinations(self):
        sheet = self.workbook["Niveau x profiel"]
        rijen = [
            [cell for cell in row]
            for row in sheet.iter_rows(min_row=3, values_only=True)
            if any(value is not None for value in row)
        ]
        kop = rijen[0]
        self.assertEqual(kop[0], "Aantallen")
        self.assertEqual(kop[-1], "Totaal")
        mbo = next(rij for rij in rijen if rij[0] == "MBO 4")
        self.assertEqual(mbo[kop.index("Techniek")], 2)
        self.assertEqual(mbo[kop.index("Zorg")], 1)
        self.assertEqual(mbo[-1], 3)

    def test_the_margins_add_up(self):
        sheet = self.workbook["Niveau x profiel"]
        totaal = next(
            row for row in sheet.iter_rows(min_row=3, values_only=True)
            if row and row[0] == "Totaal"
        )
        self.assertEqual(totaal[-1], 4)

    def test_merged_spellings_are_written_down(self):
        sheet = self.workbook["Niveau x profiel"]
        tekst = " ".join(
            str(value) for row in sheet.iter_rows(values_only=True)
            for value in row if value is not None
        )
        self.assertIn("Samengevoegde schrijfwijzen", tekst)
        self.assertIn("mbo-4", tekst)

    def test_a_workbook_without_a_crosstab_still_works(self):
        import openpyxl
        import tempfile
        from bezoekerslijst_core import export_statistics_workbook

        output = Path(tempfile.mkdtemp()) / "zonder.xlsx"
        export_statistics_workbook([("Opleidingsniveau", [("MBO", 1, 1, 0, 0)])], output)
        self.assertNotIn("Niveau x profiel", openpyxl.load_workbook(output).sheetnames)


if __name__ == "__main__":
    unittest.main()
